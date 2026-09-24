#!/usr/bin/env python3
"""Parse and safely reconcile the local Fuego Cero XLSX against matafuegos.json.

Dry-run is the default. Every mode requires explicit local input and its SHA-256.
Apply and rollback only write a distinct output path; no network code is used.
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import hashlib
import json
import os
import sys
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

VERSION = "fuego-cero-import-v1"
PROVIDER = "Fuego Cero"
PENDING_BRANCH = "147"
EXPECTED_NUMBERED_BRANCHES = 35
EXPECTED_TOTAL_EQUIPMENT = 509
HEADERS = (
    "LOCAL", "NOMBRE FANTASIA", None, "DIRECCION", "Vencimiento",
    "ABC x 5", "ABC x 10", "CO2 x 3.5", "CO2 x 5", "HCFC X 2.5", "HCFC X 5",
    "HORARIOS", "LOCALIDAD", "PROVINCIA", "SHOPPING", "PROPIETARIO",
)
EQUIPMENT_COLUMNS = {
    5: ("ABC", "5 KG", "ABC x 5"),
    6: ("ABC", "10 KG", "ABC x 10"),
    7: ("CO2", "3.5 KG", "CO2 x 3.5"),
    8: ("CO2", "5 KG", "CO2 x 5"),
    9: ("HCFC", "2.5 KG", "HCFC x 2.5"),
    10: ("HCFC", "5 KG", "HCFC x 5"),
}
OPERATIONAL_KEYS = {
    "fecha_vencimiento_manual", "historial", "historial_mantenimientos", "estado_manual",
    "actualizado_por_sucursal", "actualizado_at", "observacion_mantenimiento",
    "fecha_vencimiento_anterior", "fecha_vencimiento_original", "visita_id",
}


def canonical_bytes(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()


def object_hash(value: Any) -> str:
    return hashlib.sha256(canonical_bytes(value)).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def atomic_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(value, f, ensure_ascii=False, indent=2)
            f.write("\n")
            f.flush(); os.fsync(f.fileno())
        os.replace(name, path)
    except BaseException:
        try: os.unlink(name)
        except FileNotFoundError: pass
        raise


def normalize_branch(value: Any) -> str:
    if isinstance(value, bool):
        raise ValueError("LOCAL booleano inválido")
    text = str(value or "").strip()
    if text.endswith(".0"): text = text[:-2]
    if not text.isdigit() or not 1 <= int(text) <= 999:
        raise ValueError(f"LOCAL inválido: {value!r}")
    return f"{int(text):03d}"


def month_date(value: Any) -> str:
    if isinstance(value, (dt.datetime, dt.date)):
        return dt.date(value.year, value.month, 1).isoformat()
    text = str(value or "").strip()
    for fmt in ("%m/%Y", "%m/%y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            parsed = dt.datetime.strptime(text, fmt)
            return dt.date(parsed.year, parsed.month, 1).isoformat()
        except ValueError:
            pass
    raise ValueError(f"vencimiento mensual inválido: {value!r}")


def count_cell(value: Any, row: int, header: str) -> int:
    if value in (None, ""): return 0
    if isinstance(value, bool): raise ValueError(f"fila {row}: {header} inválido")
    try: number = int(value)
    except (TypeError, ValueError): raise ValueError(f"fila {row}: {header} debe ser entero")
    if number < 0 or float(value) != number:
        raise ValueError(f"fila {row}: {header} debe ser entero no negativo")
    return number


def parse_xlsx(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if wb.sheetnames != ["Hoja1"]:
            raise ValueError(f"hojas inesperadas: {wb.sheetnames!r}")
        ws = wb["Hoja1"]
        rows = ws.iter_rows(values_only=True)
        actual = tuple(next(rows)[:16])
        if actual != HEADERS:
            raise ValueError(f"encabezados no coinciden exactamente: {actual!r}")
        branches: dict[str, dict[str, Any]] = {}
        unnumbered: list[dict[str, Any]] = []
        total = 0
        for row_number, row in enumerate(rows, 2):
            values = tuple(row) + (None,) * max(0, 16 - len(row))
            counts = {spec[2]: count_cell(values[col], row_number, spec[2]) for col, spec in EQUIPMENT_COLUMNS.items()}
            row_total = sum(counts.values()); total += row_total
            local = values[0]
            if local in (None, ""):
                if any(v is not None and str(v).strip() for v in values[:16]):
                    name = " ".join(str(v).strip() for v in (values[1], values[2]) if v not in (None, "")).strip()
                    unnumbered.append({"row": row_number, "name": name, "equipment": row_total, "counts": counts, "reason": "LOCAL ausente; no importable automáticamente"})
                continue
            branch = normalize_branch(local)
            if branch in branches: raise ValueError(f"LOCAL duplicado: {branch}")
            due = month_date(values[4])
            equipment = []
            for col, (kind, capacity, label) in EQUIPMENT_COLUMNS.items():
                equipment.extend({"tipo": kind, "capacidad": capacity, "label": label, "ordinal": n + 1} for n in range(counts[label]))
            branches[branch] = {
                "row": row_number, "sucursal_num": branch, "sucursal": f"Sucursal {branch}",
                "local_nombre": " ".join(str(v).strip() for v in (values[1], values[2]) if v not in (None, "")).strip(),
                "fecha_vencimiento_proveedor": due, "equipment": equipment, "equipment_count": row_total,
            }
    finally:
        wb.close()
    names = " ".join(item["name"].lower() for item in unnumbered)
    if "don torcuato" not in names or "garin" not in names:
        raise ValueError("no se identificaron ambas filas sin LOCAL: Don Torcuato y Garín")
    numbered_total = sum(v["equipment_count"] for v in branches.values())
    pending_total = branches.get(PENDING_BRANCH, {}).get("equipment_count", 0)
    confirmed = sorted(set(branches) - {PENDING_BRANCH})
    return {
        "schema": "tecman.fuego-cero-source/v1", "xlsx": str(path), "xlsx_sha256": sha256_file(path),
        "branches": branches, "unnumbered": unnumbered,
        "summary": {
            "numbered_branches_observed": len(branches), "numbered_branches_expected": EXPECTED_NUMBERED_BRANCHES,
            "confirmed_branches": len(confirmed), "pending_branches": 1 if PENDING_BRANCH in branches else 0,
            "pending_branch": PENDING_BRANCH, "numbered_equipment": numbered_total,
            "pending_equipment": pending_total, "confirmed_equipment": numbered_total - pending_total,
            "unnumbered_equipment": sum(x["equipment"] for x in unnumbered), "total_equipment": total,
            "total_equipment_expected": EXPECTED_TOTAL_EQUIPMENT,
            "count_validation": "ok" if len(branches) == EXPECTED_NUMBERED_BRANCHES and total == EXPECTED_TOTAL_EQUIPMENT else "mismatch",
        },
    }


def load_inventory(path: Path, expected_sha: str) -> tuple[dict[str, Any], str]:
    actual = sha256_file(path)
    if actual != expected_sha.lower(): raise ValueError(f"precondición SHA falló: esperado {expected_sha.lower()}, actual {actual}")
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict) or not isinstance(value.get("matafuegos"), list) or not all(isinstance(x, dict) for x in value["matafuegos"]):
        raise ValueError("--input debe ser matafuegos.json con array matafuegos")
    ids = [str(x.get("id")) for x in value["matafuegos"]]
    if len(ids) != len(set(ids)): raise ValueError("IDs duplicados en input")
    return value, actual


def item_key(item: dict[str, Any]) -> tuple[str, str, str] | None:
    try: branch = normalize_branch(item.get("sucursal_num") or str(item.get("sucursal", "")).split()[-1])
    except ValueError: return None
    kind = str(item.get("tipo") or "").strip().upper()
    cap = str(item.get("capacidad") or "").strip().upper().replace(",", ".")
    aliases = {"5": "5 KG", "10": "10 KG", "3.5": "3.5 KG", "2.5": "2.5 KG", "5KG": "5 KG", "10KG": "10 KG", "3.5KG": "3.5 KG", "2.5KG": "2.5 KG"}
    cap = aliases.get(cap, cap)
    return (branch, kind, cap)


def has_activity(item: dict[str, Any]) -> bool:
    return any(item.get(k) not in (None, "", [], {}) for k in OPERATIONAL_KEYS)


def supplier_owned_inactive(item: dict[str, Any]) -> bool:
    return item.get("proveedor_nombre") == PROVIDER and item.get("proveedor_origen") == "fuego_cero_xlsx" and not has_activity(item)


def deterministic_id(source_sha: str, key: tuple[str, str, str], ordinal: int, existing: set[str]) -> str:
    salt = 0
    while True:
        raw = f"{VERSION}|{source_sha}|{'|'.join(key)}|{ordinal}|{salt}".encode()
        candidate = hashlib.sha256(raw).hexdigest()[:12]
        if candidate not in existing: return candidate
        salt += 1


def desired_after(before: dict[str, Any], branch: dict[str, Any], source_sha: str) -> dict[str, Any]:
    after = copy.deepcopy(before)
    after["proveedor_nombre"] = PROVIDER
    after["proveedor_origen"] = "fuego_cero_xlsx"
    after["proveedor_fuente_sha256"] = source_sha
    after["fecha_vencimiento_proveedor"] = branch["fecha_vencimiento_proveedor"]
    if not after.get("fecha_vencimiento_manual"):
        after["fecha_vencimiento"] = branch["fecha_vencimiento_proveedor"]
    return after


def op(kind: str, index: int | None, before: Any, after: Any) -> dict[str, Any]:
    identity = {"version": VERSION, "kind": kind, "index": index, "before": object_hash(before), "after": object_hash(after)}
    return {"operation_id": object_hash(identity), "kind": kind, "original_index": index,
            "item_id": (before or after or {}).get("id"), "before_sha256": object_hash(before), "after_sha256": object_hash(after),
            "before": before, "after": after}


def make_plan(input_path: Path, xlsx_path: Path, expected_sha: str) -> dict[str, Any]:
    inventory, input_sha = load_inventory(input_path, expected_sha)
    source = parse_xlsx(xlsx_path)
    items = inventory["matafuegos"]
    by_key: dict[tuple[str, str, str], list[tuple[int, dict[str, Any]]]] = defaultdict(list)
    for index, item in enumerate(items):
        key = item_key(item)
        if key: by_key[key].append((index, item))
    operations, conflicts = [], []
    existing_ids = {str(x.get("id")) for x in items}
    confirmed = sorted(set(source["branches"]) - {PENDING_BRANCH})
    for branch_num in confirmed:
        branch = source["branches"][branch_num]
        needs = Counter((x["tipo"], x["capacidad"]) for x in branch["equipment"])
        for (kind, capacity), needed in sorted(needs.items()):
            key = (branch_num, kind, capacity)
            current = sorted(by_key.get(key, []), key=lambda pair: str(pair[1].get("id")))
            for index, before in current[:needed]:
                after = desired_after(before, branch, source["xlsx_sha256"])
                if object_hash(before) != object_hash(after): operations.append(op("update", index, before, after))
            for ordinal in range(len(current) + 1, needed + 1):
                new_id = deterministic_id(source["xlsx_sha256"], key, ordinal, existing_ids); existing_ids.add(new_id)
                before = None
                after = desired_after({"id": new_id, "sucursal": f"Sucursal {branch_num}", "sucursal_num": branch_num,
                    "local_nombre": branch["local_nombre"], "ubicacion": "", "tipo": kind, "cantidad": 1,
                    "capacidad": capacity, "fecha_carga": "", "fecha_vencimiento": "", "estado_manual": "",
                    "origen": "fuego_cero_xlsx"}, branch, source["xlsx_sha256"])
                operations.append(op("add", None, before, after))
            for index, extra in current[needed:]:
                if supplier_owned_inactive(extra): operations.append(op("delete_safe", index, extra, None))
                else: conflicts.append({"item_id": extra.get("id"), "sucursal_num": branch_num, "tipo": kind, "capacidad": capacity,
                    "reason": "registro excedente con origen/actividad incompatible; no se elimina"})
    kinds = Counter(x["kind"] for x in operations)
    final_total = len(items) + kinds["add"] - kinds["delete_safe"]
    source_summary = source["summary"]
    return {
        "schema": "tecman.fuego-cero-plan/v1", "version": VERSION, "mode": "dry-run",
        "input": str(input_path), "input_sha256": input_sha, "xlsx": str(xlsx_path), "xlsx_sha256": source["xlsx_sha256"],
        "safety": {"network": False, "production": False, "in_place": False, "pending_147_mutations": 0},
        "source_summary": source_summary,
        "confirmed_branch_numbers": confirmed, "pending": {"147": "pendiente_confirmacion"},
        "unnumbered": source["unnumbered"],
        "summary": {"adds": kinds["add"], "updates": kinds["update"], "safe_deletes": kinds["delete_safe"],
            "conflicts": len(conflicts), "operations": len(operations), "input_total": len(items), "final_total": final_total,
            "confirmed_branches": len(confirmed), "pending_branches": 1},
        "conflicts": conflicts, "operations": operations,
    }


def ensure_distinct(**paths: Path) -> None:
    seen = {}
    for label, path in paths.items():
        resolved = path.resolve()
        if resolved in seen: raise ValueError(f"rutas deben ser distintas: {seen[resolved]} y {label}")
        seen[resolved] = label


def apply_plan(input_path: Path, expected_sha: str, plan_path: Path, output: Path, journal: Path) -> dict[str, Any]:
    ensure_distinct(input=input_path, plan=plan_path, output=output, journal=journal)
    inventory, actual = load_inventory(input_path, expected_sha)
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    if plan.get("input_sha256") != actual: raise ValueError("plan no corresponde al input")
    items = copy.deepcopy(inventory["matafuegos"])
    indexes = {str(x.get("id")): i for i, x in enumerate(items)}
    deletes = set()
    for operation in plan.get("operations", []):
        kind, item_id = operation.get("kind"), str(operation.get("item_id"))
        if kind == "add":
            if item_id in indexes: raise ValueError(f"alta ya existe: {item_id}")
            if object_hash(operation.get("after")) != operation.get("after_sha256"): raise ValueError("after corrupto")
            indexes[item_id] = len(items); items.append(copy.deepcopy(operation["after"]))
        elif kind in {"update", "delete_safe"}:
            if item_id not in indexes: raise ValueError(f"registro ausente: {item_id}")
            if object_hash(items[indexes[item_id]]) != operation.get("before_sha256"): raise ValueError(f"before divergente: {item_id}")
            if kind == "update": items[indexes[item_id]] = copy.deepcopy(operation["after"])
            else: deletes.add(item_id)
        else: raise ValueError(f"operación no permitida: {kind}")
    items = [x for x in items if str(x.get("id")) not in deletes]
    result = copy.deepcopy(inventory); result["matafuegos"] = items
    guards = {"before": object_hash(inventory), "after": object_hash(result), "before_count": len(inventory["matafuegos"]), "after_count": len(items)}
    rows = [dict(copy.deepcopy(x), journal_guards=guards) for x in plan["operations"]]
    atomic_json(output, result); atomic_json(journal, rows)
    return {"applied": len(rows), "output_sha256": sha256_file(output), "journal": str(journal)}


def rollback(input_path: Path, expected_sha: str, journal: Path, output: Path) -> dict[str, Any]:
    ensure_distinct(input=input_path, journal=journal, output=output)
    inventory, _ = load_inventory(input_path, expected_sha)
    rows = json.loads(journal.read_text(encoding="utf-8"))
    if not rows: raise ValueError("journal vacío")
    guards = rows[0].get("journal_guards")
    if any(x.get("journal_guards") != guards for x in rows) or object_hash(inventory) != guards.get("after"):
        raise ValueError("precondición after global falló")
    items = copy.deepcopy(inventory["matafuegos"])
    for operation in reversed([row for row in rows if row["kind"] == "add"]):
        item_id = str(operation.get("item_id")); index = next((i for i, x in enumerate(items) if str(x.get("id")) == item_id), None)
        if index is None or object_hash(items[index]) != operation["after_sha256"]: raise ValueError("alta divergente")
        items.pop(index)
    for operation in [row for row in rows if row["kind"] == "update"]:
        item_id = str(operation.get("item_id")); index = next((i for i, x in enumerate(items) if str(x.get("id")) == item_id), None)
        if index is None or object_hash(items[index]) != operation["after_sha256"]: raise ValueError("update divergente")
        items[index] = copy.deepcopy(operation["before"])
    for operation in sorted((row for row in rows if row["kind"] == "delete_safe"), key=lambda row: row["original_index"]):
        item_id = str(operation.get("item_id")); index = next((i for i, x in enumerate(items) if str(x.get("id")) == item_id), None)
        if index is not None: raise ValueError("baja reapareció")
        items.insert(operation["original_index"], copy.deepcopy(operation["before"]))
    result = copy.deepcopy(inventory); result["matafuegos"] = items
    if object_hash(result) != guards.get("before"): raise ValueError("rollback no reconstruyó exactamente el input")
    atomic_json(output, result)
    return {"rolled_back": len(rows), "output_sha256": sha256_file(output)}


def sanitized_report(plan: dict[str, Any]) -> dict[str, Any]:
    return {k: copy.deepcopy(plan[k]) for k in ("schema", "version", "mode", "input_sha256", "xlsx_sha256", "safety", "source_summary", "confirmed_branch_numbers", "pending", "unnumbered", "summary", "conflicts")}


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--input", type=Path, required=True); p.add_argument("--expected-sha256", required=True)
    p.add_argument("--xlsx", type=Path); p.add_argument("--plan-output", type=Path); p.add_argument("--report-output", type=Path)
    mode = p.add_mutually_exclusive_group(); mode.add_argument("--apply", action="store_true"); mode.add_argument("--rollback", action="store_true")
    p.add_argument("--plan", type=Path); p.add_argument("--journal", type=Path); p.add_argument("--output", type=Path)
    a = p.parse_args()
    try:
        if a.apply:
            if not all((a.plan, a.journal, a.output)): p.error("--apply requiere --plan --journal --output")
            result = apply_plan(a.input, a.expected_sha256, a.plan, a.output, a.journal)
        elif a.rollback:
            if not all((a.journal, a.output)): p.error("--rollback requiere --journal --output")
            result = rollback(a.input, a.expected_sha256, a.journal, a.output)
        else:
            if not all((a.xlsx, a.plan_output)): p.error("dry-run requiere --xlsx --plan-output")
            ensure_distinct(input=a.input, xlsx=a.xlsx, plan=a.plan_output, **({"report": a.report_output} if a.report_output else {}))
            plan = make_plan(a.input, a.xlsx, a.expected_sha256); atomic_json(a.plan_output, plan)
            if a.report_output: atomic_json(a.report_output, sanitized_report(plan))
            result = {"plan": str(a.plan_output), "report": str(a.report_output) if a.report_output else None, "summary": plan["summary"], "source_summary": plan["source_summary"]}
        print(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True))
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr); raise SystemExit(2)


if __name__ == "__main__": main()
