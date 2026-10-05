"""Cálculo puro del resumen ejecutivo administrativo de Tecman."""
from __future__ import annotations

import datetime as dt
import io
import math
import re
import unicodedata
from collections import Counter, defaultdict
from statistics import median

FINAL_STATES = {"resuelto", "cerrado"}
TERMINAL_OTHER_STATES = {"rechazado", "cancelado", "anulado"}
MAX_RANGE_DAYS = 366
MIN_TIMING_SAMPLE = 3
SYNTHETIC_MARKERS = ("seed", "fixture", "tutorial", "demo")
SAFE_BRANCH_RE = re.compile(r"(?:sucursal\s*)?(\d{1,4})\b", re.I)
FORMULA_PREFIXES = ("=", "+", "-", "@")


def _text(value):
    return str(value or "").strip()


def _norm(value):
    value = unicodedata.normalize("NFKD", _text(value)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", value).strip()


def _date(value):
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time())
    raw = _text(value)
    if not raw:
        return None
    try:
        parsed = dt.datetime.fromisoformat(raw.replace("Z", "+00:00"))
        if parsed.tzinfo:
            parsed = parsed.astimezone(dt.timezone.utc).replace(tzinfo=None)
        return parsed
    except (TypeError, ValueError):
        try:
            return dt.datetime.strptime(raw[:10], "%Y-%m-%d")
        except (TypeError, ValueError):
            return None


def branch_code(value):
    raw = _text(value)
    match = SAFE_BRANCH_RE.fullmatch(raw) or SAFE_BRANCH_RE.search(raw)
    return match.group(1).zfill(3) if match else ""


def ticket_branch(ticket):
    return branch_code(ticket.get("sucursal_num") or ticket.get("sucursal") or ticket.get("branch_code"))


def _synthetic_reason(ticket):
    for marker in SYNTHETIC_MARKERS:
        if ticket.get(marker) is True:
            return marker
    structured = " ".join(_norm(ticket.get(key)) for key in (
        "fixture", "seed", "demo", "tutorial", "origen", "source", "marca", "mark", "tipo_registro"
    ))
    if ticket.get("es_sintetico") is True or ticket.get("synthetic") is True:
        return "marca_sintetica"
    for marker in SYNTHETIC_MARKERS:
        if re.search(rf"(?:^|[^a-z0-9]){marker}(?:[^a-z0-9]|$)", structured):
            return marker
    # Reglas textuales deliberadamente exactas; no se filtran términos vagos como "prueba" solos.
    visible = " ".join(_norm(ticket.get(key)) for key in ("descripcion", "titulo", "asunto"))
    if "ticket prueba" in visible:
        return "ticket_prueba"
    if re.search(r"(?:^|[^a-z0-9])demo(?:[^a-z0-9]|$)", visible):
        return "demo"
    return ""


def parse_period(args, tickets):
    desde_raw = _text(args.get("desde"))
    hasta_raw = _text(args.get("hasta"))
    if bool(desde_raw) != bool(hasta_raw):
        raise ValueError("desde y hasta deben indicarse juntos con formato YYYY-MM-DD")
    if desde_raw:
        try:
            desde = dt.datetime.strptime(desde_raw, "%Y-%m-%d")
            hasta = dt.datetime.strptime(hasta_raw, "%Y-%m-%d")
        except ValueError as exc:
            raise ValueError("desde y hasta deben usar el formato YYYY-MM-DD") from exc
        days = (hasta - desde).days
        if days < 1 or days > MAX_RANGE_DAYS:
            raise ValueError(f"El período debe abarcar entre 1 y {MAX_RANGE_DAYS} días; hasta es exclusivo")
        return desde, hasta, "consulta"
    candidates = []
    for ticket in tickets:
        if not _synthetic_reason(ticket):
            created = _date(ticket.get("creado") or ticket.get("fecha_creacion"))
            if created:
                candidates.append(created)
    if not candidates:
        return None, None, "automatico_sin_datos"
    desde = min(candidates).replace(hour=0, minute=0, second=0, microsecond=0)
    return desde, desde + dt.timedelta(days=30), "automatico"


def _counter_rows(counter, total=None, limit=None):
    rows = []
    values = sorted(counter.items(), key=lambda item: (-item[1], item[0].casefold()))
    if limit is not None:
        values = values[:limit]
    for label, count in values:
        row = {"nombre": label or "Sin dato", "cantidad": count}
        if total is not None:
            row["porcentaje"] = round(count * 100 / total, 1) if total else 0.0
        rows.append(row)
    return rows


def _percentile(values, percentile):
    if not values:
        return None
    ordered = sorted(values)
    rank = (len(ordered) - 1) * percentile
    low, high = math.floor(rank), math.ceil(rank)
    if low == high:
        return round(ordered[low], 2)
    return round(ordered[low] + (ordered[high] - ordered[low]) * (rank - low), 2)


def _timing(values):
    if len(values) < MIN_TIMING_SAMPLE:
        return {"estado": "sin datos suficientes", "muestra": len(values), "mediana_horas": None, "p90_horas": None}
    return {
        "estado": "disponible", "muestra": len(values),
        "mediana_horas": round(median(values), 2), "p90_horas": _percentile(values, .9),
    }


def _history_dates(ticket):
    result = []
    for key in ("historial", "notas", "respuestas"):
        values = ticket.get(key) or []
        if isinstance(values, list):
            for item in values:
                if isinstance(item, dict):
                    value = _date(item.get("fecha") or item.get("created_at") or item.get("timestamp"))
                    if value:
                        result.append(value)
    # ``actualizado`` nace igual a ``creado`` y no prueba una gestión humana.
    for key in ("primera_gestion_at", "fecha_primera_gestion"):
        value = _date(ticket.get(key))
        if value:
            result.append(value)
    return result


def _resolution_date(ticket):
    for key in ("fecha_cierre", "resuelto_at", "cerrado_at", "fecha_resolucion"):
        value = _date(ticket.get(key))
        if value:
            return value
    return None


def _is_material_purchase(ticket):
    values = " ".join(_norm(ticket.get(k)) for k in ("categoria", "subcategoria", "tipo", "sector", "asignado"))
    return any(term in values for term in ("material", "insumo", "compras"))


def _purchase_derivation_in_period(ticket, desde, hasta):
    """Detecta sólo derivaciones estructuradas y fechadas dentro del período."""
    requisition = ticket.get("compra_np_requisicion")
    if isinstance(requisition, dict):
        created = _date(requisition.get("creada") or requisition.get("created_at"))
        if created and desde and hasta and desde <= created < hasta:
            return True
    for key in ("historial", "notas"):
        for item in ticket.get(key) or []:
            if not isinstance(item, dict) or item.get("tipo") != "compra_np_derivado_compras":
                continue
            happened = _date(item.get("fecha") or item.get("created_at"))
            if happened and desde and hasta and desde <= happened < hasta:
                return True
    return False


def _mandatory_attachment_missing(ticket):
    required = ticket.get("adjunto_obligatorio") is True or ticket.get("requiere_adjunto") is True
    if not required:
        return False
    attachments = []
    for key in ("adjuntos", "fotos", "archivos"):
        value = ticket.get(key)
        if isinstance(value, list):
            attachments.extend(value)
        elif value:
            attachments.append(value)
    return not bool(attachments)


def _event_dict(event):
    if isinstance(event, dict):
        return event
    return {
        "user_id": getattr(event, "user_id", None), "event_type": getattr(event, "event_type", None),
        "provider": getattr(event, "provider", None), "details": getattr(event, "details", {}) or {},
        "created_at": getattr(event, "created_at", None),
    }


def _user_dict(user):
    if isinstance(user, dict):
        return user
    return {
        "id": getattr(user, "id", None), "role": getattr(user, "role", None),
        "sucursal_num": getattr(user, "sucursal_num", None), "sucursal": getattr(user, "sucursal", None),
        "email": getattr(user, "email", None),
    }


def rank_sucursal_logins(events, desde=None, hasta=None):
    """Rankea logins atribuibles usando sólo metadata canónica y no sensible."""
    counts = Counter()
    labels = {}
    for raw in events or []:
        event = _event_dict(raw)
        if event.get("event_type") != "login_success":
            continue
        details = event.get("details") if isinstance(event.get("details"), dict) else {}
        if details.get("role") != "sucursal":
            continue
        happened = _date(event.get("created_at") or event.get("fecha"))
        if desde and (not happened or happened < desde):
            continue
        if hasta and (not happened or happened >= hasta):
            continue
        raw_code = _text(details.get("sucursal_num"))
        raw_label = _text(details.get("sucursal_label"))
        code = branch_code(raw_code)
        if not code and raw_code.casefold() in {"central", "garin"}:
            code = raw_code.casefold()
        if not code or not raw_label:
            continue
        counts[code] += 1
        labels.setdefault(code, raw_label)
    return [
        {"sucursal_num": code, "sucursal_label": labels[code], "ingresos": count}
        for code, count in sorted(counts.items(), key=lambda item: (-item[1], item[0]))
    ]


def login_metrics(events, users, catalog_codes, desde, hasta, branch_email_map=None):
    branch_email_map = {
        _text(email).casefold(): branch_code(code)
        for email, code in (branch_email_map or {}).items()
        if _text(email) and branch_code(code) in catalog_codes
    }
    user_branch = {}
    for raw in users or []:
        user = _user_dict(raw)
        code = branch_code(user.get("sucursal_num") or user.get("sucursal") or user.get("branch_code"))
        if not code:
            code = branch_email_map.get(_text(user.get("email")).casefold(), "")
        if code in catalog_codes and user.get("id"):
            user_branch[str(user["id"])] = code
    success = []
    mapped = Counter()
    for raw in events or []:
        event = _event_dict(raw)
        if event.get("event_type") != "login_success":
            continue
        happened = _date(event.get("created_at") or event.get("fecha"))
        if not happened:
            continue
        details = event.get("details") if isinstance(event.get("details"), dict) else {}
        code = branch_code(details.get("sucursal_num") or details.get("sucursal") or details.get("branch_code"))
        if not code:
            code = user_branch.get(str(event.get("user_id") or ""), "")
        success.append((happened, code))
        if desde and hasta and desde <= happened < hasta and code in catalog_codes:
            mapped[code] += 1
    dates = [value[0] for value in success]
    all_attributable = bool(success) and all(code in catalog_codes for _, code in success)
    sufficient = bool(desde and hasta and dates and min(dates) <= desde and max(dates) >= hasta and all_attributable)
    with_login = set(mapped)
    missing = sorted(catalog_codes - with_login)
    return {
        "eventos_login_success_periodo": sum(mapped.values()),
        "eventos_login_success_totales": len(success),
        "sucursales_con_evidencia": len(with_login),
        "sucursales_sin_evidencia": len(missing),
        "sin_evidencia_codigos": missing,
        "etiqueta_ausencia": "nunca ingresaron" if sufficient else "sin evidencia de ingreso",
        "cobertura_desde": min(dates).isoformat() if dates else None,
        "cobertura_hasta": max(dates).isoformat() if dates else None,
        "atribucion_completa": all_attributable,
        "cobertura_suficiente": sufficient,
        "ranking_sucursales": rank_sucursal_logins(events, desde, hasta),
    }


def build_report(tickets, catalog, auth_events=None, users=None, query=None, now=None, branch_email_map=None):
    tickets = list(tickets or [])
    catalog = catalog or {}
    catalog_codes = {str(code).zfill(3) for code in catalog}
    desde, hasta, period_source = parse_period(query or {}, tickets)
    excluded = Counter()
    selected = []
    for ticket in tickets:
        reason = _synthetic_reason(ticket)
        if reason:
            excluded[reason] += 1
            continue
        created = _date(ticket.get("creado") or ticket.get("fecha_creacion"))
        if desde and hasta and created and desde <= created < hasta:
            selected.append((ticket, created))
    total = len(selected)
    daily, weekly, branches, categories, subcategories, types, states = (Counter() for _ in range(7))
    responsibles, sectors = Counter(), Counter()
    pending_responsibles, pending_sectors = Counter(), Counter()
    first_times, resolution_times = [], []
    pending = []
    quality = Counter()
    now = now or dt.datetime.now()
    for ticket, created in selected:
        daily[created.date().isoformat()] += 1
        monday = (created.date() - dt.timedelta(days=created.weekday())).isoformat()
        weekly[monday] += 1
        code = ticket_branch(ticket)
        branches[code or "No normalizada"] += 1
        categories[_text(ticket.get("categoria")) or "Sin dato"] += 1
        subcategories[_text(ticket.get("subcategoria")) or "Sin dato"] += 1
        types[_text(ticket.get("tipo")) or "Sin dato"] += 1
        state = _text(ticket.get("estado")) or "Sin dato"
        states[state] += 1
        responsibles[_text(ticket.get("asignado") or ticket.get("responsable")) or "Sin asignar"] += 1
        sectors[_text(ticket.get("sector")) or "Sin sector"] += 1
        if any(not _text(ticket.get(key)) for key in ("creado", "estado", "categoria", "subcategoria", "descripcion")) or not code:
            quality["obligatorios_vacios"] += 1
        if not code or code not in catalog_codes:
            quality["sucursal_no_normalizada"] += 1
        if _mandatory_attachment_missing(ticket):
            quality["adjunto_obligatorio_faltante"] += 1
        history_dates = [value for value in _history_dates(ticket) if value >= created]
        if history_dates:
            first_times.append((min(history_dates) - created).total_seconds() / 3600)
        resolved = _resolution_date(ticket)
        if resolved and resolved >= created and _norm(state) in FINAL_STATES:
            resolution_times.append((resolved - created).total_seconds() / 3600)
        if _norm(state) not in FINAL_STATES | TERMINAL_OTHER_STATES:
            age = max(0, (now - created).days)
            pending_responsibles[_text(ticket.get("asignado") or ticket.get("responsable")) or "Sin asignar"] += 1
            pending_sectors[_text(ticket.get("sector")) or "Sin sector"] += 1
            pending.append({
                "sucursal": code or "No normalizada", "estado": state,
                "antiguedad_dias": age, "responsable": _text(ticket.get("asignado") or ticket.get("responsable")) or "Sin asignar",
                "sector": _text(ticket.get("sector")) or "Sin sector", "categoria": _text(ticket.get("categoria")) or "Sin dato",
                "compras_insumos": "Sí" if _is_material_purchase(ticket) else "No",
            })
    branch_with = {code for code in branches if code in catalog_codes}
    branch_without = sorted(catalog_codes - branch_with)
    branch_rows = []
    for code in sorted(catalog_codes):
        info = catalog.get(code, {}) or {}
        branch_rows.append({
            "codigo": code, "nombre": _text(info.get("tienda") or info.get("nombre")) or f"Sucursal {code}",
            "tickets": branches.get(code, 0), "situacion": "Con tickets" if branches.get(code) else "Sin tickets",
        })
    for code in sorted(set(branches) - catalog_codes):
        branch_rows.append({"codigo": code, "nombre": "No catalogada", "tickets": branches[code], "situacion": "Fuera de catálogo"})
    top = _counter_rows(branches, total=total, limit=10)
    top5_count = sum(row["cantidad"] for row in top[:5])
    top10_count = sum(row["cantidad"] for row in top)
    final_count = sum(count for state, count in states.items() if _norm(state) in FINAL_STATES)
    backlog_ages = [row["antiguedad_dias"] for row in pending]
    purchase_count = sum(1 for ticket, _ in selected if _is_material_purchase(ticket))
    purchase_derivations = sum(
        1 for ticket in tickets
        if not _synthetic_reason(ticket) and _purchase_derivation_in_period(ticket, desde, hasta)
    )
    timings = {"primera_gestion": _timing(first_times), "resolucion": _timing(resolution_times)}
    backlog_ratio = len(pending) / total if total else 0
    aged_ratio = sum(age >= 30 for age in backlog_ages) / len(backlog_ages) if backlog_ages else 0
    purchase_ratio = purchase_count / total if total else 0
    adoption_alert = len(branch_without) >= 5 or bool(catalog_codes and len(branch_without) / len(catalog_codes) >= .20)
    recommendations = [
        (
            "Validar disponibilidad y período de las fuentes antes de interpretar resultados: el volumen incluido es cero."
            if total == 0 else
            "Conservar este corte como línea de base y comparar el siguiente período con las mismas reglas: el volumen incluido es mayor que cero."
        ),
        (
            "Revisar causas en las cinco sucursales de mayor volumen: concentran al menos 60% de los tickets."
            if total and top5_count / total >= .60 else
            "Mantener seguimiento mensual de concentración: el top 5 está por debajo del umbral de 60%."
        ),
        (
            "Priorizar una revisión semanal del backlog: los pendientes representan al menos 30% del período."
            if backlog_ratio >= .30 else
            "Mantener el control periódico del backlog: permanece por debajo del umbral de 30% del período."
        ),
        (
            "Armar un plan de cierre para pendientes de 30 días o más: representan al menos 20% del backlog."
            if aged_ratio >= .20 else
            "Mantener la priorización vigente: los pendientes de 30 días o más están por debajo de 20% del backlog."
        ),
        (
            "Corregir captura y normalización: existe al menos un ticket con campos obligatorios vacíos o sucursal no normalizada."
            if quality["obligatorios_vacios"] or quality["sucursal_no_normalizada"] else
            "Sostener los controles de captura: no se detectaron obligatorios vacíos ni sucursales no normalizadas."
        ),
        (
            "Separar el seguimiento de Compras/insumos: representa al menos 20% del volumen del período."
            if purchase_ratio >= .20 else
            "Mantener Compras/insumos dentro del seguimiento general: representa menos de 20% del volumen."
        ),
        (
            "Completar historial estructurado de gestión y cierre: alguna métrica temporal tiene menos de tres tickets."
            if any(item["estado"] != "disponible" for item in timings.values()) else
            "Monitorear mediana y P90: ambas métricas temporales alcanzan la muestra mínima de tres tickets."
        ),
        (
            "Contrastar adopción con sucursales sin tickets: alcanzan al menos cinco o 20% del catálogo."
            if adoption_alert else
            "Mantener seguimiento de adopción: las sucursales sin tickets están por debajo de cinco y de 20% del catálogo."
        ),
    ]
    login = login_metrics(auth_events or [], users or [], catalog_codes, desde, hasta, branch_email_map=branch_email_map)
    report = {
        "periodo": {
            "desde": desde.date().isoformat() if desde else None, "hasta_exclusivo": hasta.date().isoformat() if hasta else None,
            "fuente": period_source, "dias": (hasta - desde).days if desde and hasta else 0,
        },
        "metodologia": {
            "intervalo": "[desde, hasta)", "limite_consulta_dias": MAX_RANGE_DAYS,
            "exclusiones": dict(excluded), "excluidos_total": sum(excluded.values()),
            "reglas_exclusion": ["marcas estructuradas seed/fixture/tutorial/demo", "bandera sintética", "frase exacta TICKET PRUEBA"],
            "universo_sucursales": "Catálogo disponible; no se rotula habilitadas porque la fuente no informa estado activo.",
            "backlog": "Estado actual de tickets creados en el período; no es snapshot histórico al día 30.",
            "tiempos": f"Mediana/P90 sólo con al menos {MIN_TIMING_SAMPLE} tickets y marcas estructuradas no negativas.",
            "umbrales_recomendaciones": [
                "volumen = 0", "concentración top 5 ≥ 60%", "backlog ≥ 30% del volumen",
                "pendientes de 30+ días ≥ 20% del backlog", "calidad con al menos 1 hallazgo",
                "Compras/insumos ≥ 20%", "tiempos con muestra mínima de 3",
                "sin tickets ≥ 5 sucursales o ≥ 20% del catálogo",
            ],
        },
        "metricas": {
            "total": total, "promedio_diario": round(total / ((hasta - desde).days if desde and hasta else 1), 2),
            "promedio_semanal": round(total / (((hasta - desde).days / 7) if desde and hasta else 1), 2),
            "finalizados": final_count, "tasa_finalizacion": round(final_count * 100 / total, 1) if total else 0.0,
            "backlog": len(pending), "backlog_7_mas": sum(age >= 7 for age in backlog_ages),
            "backlog_30_mas": sum(age >= 30 for age in backlog_ages), "compras_insumos": purchase_count,
            "derivaciones_compras_insumos": purchase_derivations,
            "concentracion_top5": round(top5_count * 100 / total, 1) if total else 0.0,
            "concentracion_top10": round(top10_count * 100 / total, 1) if total else 0.0,
        },
        "evolucion_diaria": [{"fecha": key, "cantidad": daily[key]} for key in sorted(daily)],
        "evolucion_semanal": [{"semana": key, "cantidad": weekly[key]} for key in sorted(weekly)],
        "top_sucursales": top, "sucursales": branch_rows, "sucursales_distribucion": _counter_rows(branches, total),
        "sucursales_resumen": {"catalogadas": len(catalog_codes), "con_tickets": len(branch_with), "sin_tickets": len(branch_without)},
        "categorias": _counter_rows(categories, total), "subcategorias": _counter_rows(subcategories, total),
        "tipos": _counter_rows(types, total), "estados": _counter_rows(states, total),
        "responsables": _counter_rows(responsibles, total), "sectores": _counter_rows(sectors, total),
        "responsables_carga": _counter_rows(pending_responsibles, len(pending)),
        "sectores_carga": _counter_rows(pending_sectors, len(pending)),
        "tiempos": timings, "pendientes": sorted(pending, key=lambda row: (-row["antiguedad_dias"], row["sucursal"])),
        "calidad": dict(quality), "login": login, "recomendaciones": recommendations,
    }
    reconcile(report)
    return report


def reconcile(report):
    total = report["metricas"]["total"]
    checks = {
        "diario": sum(row["cantidad"] for row in report["evolucion_diaria"]),
        "semanal": sum(row["cantidad"] for row in report["evolucion_semanal"]),
        "categorias": sum(row["cantidad"] for row in report["categorias"]),
        "estados": sum(row["cantidad"] for row in report["estados"]),
        "sucursales": sum(row["cantidad"] for row in report["sucursales_distribucion"]),
    }
    if any(value != total for value in checks.values()):
        raise ValueError(f"Reconciliación de totales fallida: {checks}, total={total}")
    summary = report["sucursales_resumen"]
    if summary["con_tickets"] + summary["sin_tickets"] != summary["catalogadas"]:
        raise ValueError("Reconciliación del catálogo de sucursales fallida")
    top5 = sum(row["cantidad"] for row in report["top_sucursales"][:5])
    top10 = sum(row["cantidad"] for row in report["top_sucursales"][:10])
    if not (top5 <= top10 <= total):
        raise ValueError("Reconciliación de concentración fallida")
    for key in ("primera_gestion", "resolucion"):
        values = report["tiempos"][key]
        if any(value is not None and value < 0 for value in (values["mediana_horas"], values["p90_horas"])):
            raise ValueError("Se detectaron tiempos negativos")
    report["reconciliaciones"] = {**{key: value == total for key, value in checks.items()}, "catalogo": True, "concentracion": True, "tiempos_no_negativos": True}


def xlsx_safe(value):
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value)
    if text.startswith(FORMULA_PREFIXES) or text.startswith(("\t", "\r", "\n")):
        return "'" + text
    return text


def build_workbook(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    sheets = ["Resumen", "Sucursales", "Categorias", "Estados", "Pendientes"]
    for name in sheets:
        wb.create_sheet(name)
    header_fill = PatternFill("solid", fgColor="DCE6F1")

    def add_table(ws, headers, rows, start=1):
        for col, value in enumerate(headers, 1):
            cell = ws.cell(start, col, xlsx_safe(value)); cell.font = Font(bold=True); cell.fill = header_fill
        for ridx, row in enumerate(rows, start + 1):
            for col, value in enumerate(row, 1):
                ws.cell(ridx, col, xlsx_safe(value))
        ws.freeze_panes = ws.cell(start + 1, 1)
        ws.auto_filter.ref = f"A{start}:{chr(64 + min(len(headers), 26))}{max(start, start + len(rows))}"
        for column in ws.columns:
            letter = column[0].column_letter
            ws.column_dimensions[letter].width = min(42, max(10, max(len(str(cell.value or "")) for cell in column) + 2))

    ws = wb["Resumen"]
    summary = [
        ("Resumen ejecutivo TECMAN", ""), ("Desde", report["periodo"]["desde"] or "Sin datos"),
        ("Hasta (exclusivo)", report["periodo"]["hasta_exclusivo"] or "Sin datos"),
        ("Tickets", report["metricas"]["total"]), ("Promedio diario", report["metricas"]["promedio_diario"]),
        ("Promedio semanal", report["metricas"]["promedio_semanal"]), ("Finalizados", report["metricas"]["finalizados"]),
        ("Tasa finalización (%)", report["metricas"]["tasa_finalizacion"]), ("Backlog", report["metricas"]["backlog"]),
        ("Derivaciones a Compras/insumos", report["metricas"]["derivaciones_compras_insumos"]),
        ("Top 5 (%)", report["metricas"]["concentracion_top5"]), ("Top 10 (%)", report["metricas"]["concentracion_top10"]),
        ("Excluidos sintéticos", report["metodologia"]["excluidos_total"]),
        ("Cobertura login", report["login"]["etiqueta_ausencia"]),
        ("Primera gestión", report["tiempos"]["primera_gestion"]["estado"]),
        ("Resolución", report["tiempos"]["resolucion"]["estado"]),
    ]
    add_table(ws, ["Indicador", "Valor"], summary)
    start = len(summary) + 4
    add_table(ws, ["Recomendaciones (separadas de hechos)"], [(item,) for item in report["recomendaciones"]], start=start)
    add_table(wb["Sucursales"], ["Código", "Sucursal", "Tickets", "Situación"], [(r["codigo"], r["nombre"], r["tickets"], r["situacion"]) for r in report["sucursales"]])
    category_rows = []
    for level, key in (("Categoría", "categorias"), ("Subcategoría", "subcategorias"), ("Tipo", "tipos")):
        category_rows.extend((level, row["nombre"], row["cantidad"], row["porcentaje"]) for row in report[key])
    add_table(wb["Categorias"], ["Nivel", "Nombre", "Cantidad", "%"], category_rows)
    add_table(wb["Estados"], ["Estado", "Cantidad", "%"], [(r["nombre"], r["cantidad"], r["porcentaje"]) for r in report["estados"]])
    add_table(wb["Pendientes"], ["Sucursal", "Estado", "Antigüedad (días)", "Responsable", "Sector", "Categoría", "Compras/insumos"], [tuple(row[key] for key in ("sucursal", "estado", "antiguedad_dias", "responsable", "sector", "categoria", "compras_insumos")) for row in report["pendientes"]])
    return wb


def workbook_bytes(report):
    wb = build_workbook(report)
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    output.seek(0)
    return output
