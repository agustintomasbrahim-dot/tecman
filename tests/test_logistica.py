import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook

# Aislar todos los efectos de importación de app.py antes de importarlo.
_TMP = tempfile.TemporaryDirectory()
os.environ["TECMAN_DATA_DIR"] = str(Path(_TMP.name) / "data")
os.environ["TECMAN_UPLOADS_DIR"] = str(Path(_TMP.name) / "uploads")
os.environ.pop("DATABASE_URL", None)
os.environ["LOGISTICA_INSUMOS_SUCURSALES_ENABLED"] = "false"

import app as tecman  # noqa: E402
from logistica import LogisticsError, LogisticsService, LogisticsStore, logistics_entra_role  # noqa: E402


class LogisticsServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "logistica.json"
        self.service = LogisticsService(LogisticsStore(self.path))

    def tearDown(self):
        self.temp.cleanup()

    def stock(self, csv_text="item,cantidad\nGuantes,5\n"):
        return self.service.import_stock(csv_text.encode(), "stock.csv", "Dabra")

    def order(self, branch="Sucursal 011", item="Guantes", qty=4):
        return self.service.create_order(branch, [{"item": item, "requested": qty}], "Sucursal")

    def delivered_order(self, branch="Sucursal 011"):
        self.stock()
        order = self.order(branch)
        self.service.approve(order["id"], {"Guantes": 4}, "Dabra")
        wave = self.service.create_wave([order["id"]], "semanal", "Dabra")
        self.service.update_wave(wave["id"], "retirado", "R1", "", "Garín")
        self.service.update_wave(wave["id"], "en_distribucion", "R1", "", "Garín")
        self.service.update_wave(wave["id"], "entregado", "R1", "", "Garín")
        return order, wave

    def test_reserva_no_sobreasigna_y_faltante_notifica_compras(self):
        self.stock()
        first = self.order(qty=4)
        second = self.order(branch="Sucursal 014", qty=4)
        self.service.approve(first["id"], {"Guantes": 4}, "Dabra")
        result = self.service.approve(second["id"], {"Guantes": 4}, "Dabra")
        self.assertEqual(result["lines"][0]["reserved"], 1)
        self.assertEqual(result["lines"][0]["shortage"], 3)
        self.assertEqual(sum(o["lines"][0]["reserved"] for o in self.service.state()["orders"].values()), 5)
        self.assertTrue(any(n["audience"] == "compras" and n["order_id"] == second["id"]
                            for n in self.service.state()["notifications"]))

    def test_stock_suficiente_reserva_y_deja_preparado_para_garin(self):
        self.stock()
        order = self.order(qty=4)
        result = self.service.approve(order["id"], {"Guantes": 4}, "Dabra")
        self.assertEqual(result["status"], "preparado_retiro_garin")
        self.assertEqual(result["fulfillment_status"], "preparado_retiro_garin")
        self.assertEqual(result["lines"][0]["reserved"], 4)
        self.assertEqual(self.service.state().get("requisitions"), {})

    def test_faltante_parcial_y_total_crean_requisiciones_auditables(self):
        self.stock()
        partial = self.order(qty=8)
        total = self.order(branch="Sucursal 014", item="Barbijos", qty=3)
        partial_result = self.service.approve(partial["id"], {"Guantes": 8}, "Dabra")
        total_result = self.service.approve(total["id"], {"Barbijos": 3}, "Dabra")
        state = self.service.state()
        partial_req = state["requisitions"][partial_result["requisition_id"]]
        total_req = state["requisitions"][total_result["requisition_id"]]
        self.assertEqual((partial_result["lines"][0]["reserved"], partial_result["lines"][0]["shortage"]), (5, 3))
        self.assertEqual((total_result["lines"][0]["reserved"], total_result["lines"][0]["shortage"]), (0, 3))
        self.assertEqual((partial_req["number"], total_req["number"]), ("REQ-000001", "REQ-000002"))
        self.assertEqual(partial_req["status"], "pendiente_compras")
        self.assertEqual(partial_req["email_status"], "pending")

    def test_revalidacion_sin_cambios_no_duplica_requisicion_ni_email(self):
        order = self.order(qty=4)
        first = self.service.approve(order["id"], {"Guantes": 4}, "Dabra")
        sender = Mock()
        sent = self.service.send_requisition_email(first["requisition_id"], sender, "Dabra")
        second = self.service.approve(order["id"], {"Guantes": 4}, "Dabra")
        duplicate = self.service.send_requisition_email(second["requisition_id"], sender, "Dabra")
        self.assertEqual(sent["status"], "sent")
        self.assertEqual(duplicate["status"], "duplicate")
        self.assertEqual(first["requisition_id"], second["requisition_id"])
        self.assertEqual(len(self.service.state()["requisitions"]), 1)
        sender.assert_called_once()

    def test_falla_smtp_queda_visible_y_reintento_es_seguro(self):
        order = self.order(qty=2)
        approved = self.service.approve(order["id"], {"Guantes": 2}, "Dabra")
        failing = Mock(side_effect=RuntimeError("SMTP demo caído"))
        failed = self.service.send_requisition_email(approved["requisition_id"], failing, "Dabra")
        skipped = self.service.send_requisition_email(approved["requisition_id"], failing, "Dabra")
        successful = Mock()
        retried = self.service.send_requisition_email(approved["requisition_id"], successful, "Dabra", retry=True)
        duplicate = self.service.send_requisition_email(approved["requisition_id"], successful, "Dabra", retry=True)
        self.assertEqual((failed["status"], skipped["status"], retried["status"], duplicate["status"]),
                         ("error", "duplicate", "sent", "duplicate"))
        req = self.service.state()["requisitions"][approved["requisition_id"]]
        self.assertEqual(req["email_status"], "sent")
        self.assertEqual(req["email_attempts"], 2)
        failing.assert_called_once()
        successful.assert_called_once()

    def test_importacion_csv_y_xlsx_es_atomica(self):
        self.stock()
        before = self.service.state()["stock"]
        with self.assertRaises(LogisticsError):
            self.service.import_stock(b"item,cantidad\nBarbijo,2\nMalo,no\n", "stock.csv", "Dabra")
        self.assertEqual(self.service.state()["stock"], before)

        template = self.service.stock_template()
        wb = load_workbook(io.BytesIO(template))
        self.assertEqual(wb.active.max_row, 1)
        wb.active.append(["Barbijo", 7])
        out = io.BytesIO(); wb.save(out)
        self.service.import_stock(out.getvalue(), "stock.xlsx", "Dabra")
        self.assertEqual(self.service.state()["stock"], {"Barbijo": 7})

    def test_importacion_no_invalida_reservas_activas(self):
        self.stock()
        order = self.order(qty=4)
        self.service.approve(order["id"], {"Guantes": 4}, "Dabra")
        before = self.service.state()["stock"]
        with self.assertRaises(LogisticsError):
            self.service.import_stock(b"item,cantidad\nGuantes,3\n", "stock.csv", "Dabra")
        self.assertEqual(self.service.state()["stock"], before)

    def test_creacion_ola_export_xlsx_y_descuento_unico(self):
        self.stock()
        order = self.order()
        self.service.approve(order["id"], {"Guantes": 4}, "Dabra")
        wave = self.service.create_wave([order["id"]], "quincenal", "Dabra")
        book = load_workbook(io.BytesIO(self.service.export_wave(wave["id"])), data_only=True)
        rows = list(book.active.iter_rows(values_only=True))
        self.assertEqual(rows[1][3:7], ("Sucursal 011", order["id"], "Guantes", 4))

        self.service.update_wave(wave["id"], "retirado", "Ruta norte", "Retiro OK", "Garín")
        self.service.update_wave(wave["id"], "retirado", "Ruta norte", "reintento", "Garín")
        state = self.service.state()
        self.assertEqual(state["stock"]["Guantes"], 1)
        self.assertEqual(len(state["movements"]), 1)
        self.assertEqual(state["movements"][0]["id"], f"ola:{wave['id']}:Guantes")

    def test_sucursal_confirma_solo_su_pedido_e_idempotente(self):
        order, _ = self.delivered_order("Sucursal 011")
        with self.assertRaises(PermissionError):
            self.service.confirm_receipt(order["id"], "Sucursal 014", "Otra")
        first = self.service.confirm_receipt(order["id"], "Sucursal 011", "Sucursal")
        second = self.service.confirm_receipt(order["id"], "Sucursal 011", "Sucursal")
        self.assertEqual(first["received_at"], second["received_at"])
        events = [h for h in self.service.state()["orders"][order["id"]]["history"] if h["action"] == "recepcion_confirmada"]
        self.assertEqual(len(events), 1)


class LogisticsRouteTests(unittest.TestCase):
    def setUp(self):
        tecman.app.config.update(
            TESTING=True,
            LOGISTICA_INSUMOS_SUCURSALES_ENABLED=False,
            LOGISTICA_PORTAL_TEST_MODE=True,
            LOGISTICA_SMTP_MOCK=False,
            LOGISTICA_LOCAL_LOGIN_ENABLED=False,
        )
        self.client = tecman.app.test_client()
        path = tecman.logistica_service.store.json_path
        if path.exists():
            path.unlink()
        tecman.app.config["LOGISTICA_LOCAL_USERS"] = {
            "dabra": {"password": "test-dabra", "role": "dabra", "name": "Dabra Test"},
            "garin": {"password": "test-garin", "role": "garin", "name": "Garín Test"},
        }
        self.smtp_mock = patch.object(tecman, "_smtp_send", autospec=True).start()
        self.addCleanup(patch.stopall)

    def session(self, **values):
        if values.get("logistica_role") == "dabra":
            values["logistica_user"] = tecman.PREPARADOR_DABRA_EMAIL
            values.setdefault("logistica_name", tecman.PREPARADOR_DABRA_NAME)
            values["auth_provider"] = "entra"
        with self.client.session_transaction() as sess:
            sess.clear(); sess.update(values); sess["_csrf_token"] = "csrf-test"

    def post(self, url, data=None):
        payload = {"_csrf_token": "csrf-test"}; payload.update(data or {})
        return self.client.post(url, data=payload)

    def test_permisos_dabra_garin_y_csrf(self):
        self.session(logistica_role="dabra", logistica_user="d")
        self.assertEqual(self.client.get("/logistica").status_code, 200)
        self.assertEqual(self.client.get("/logistica/garin").status_code, 403)
        self.assertEqual(self.client.post("/logistica/olas", data={}).status_code, 400)
        self.session(logistica_role="garin", logistica_user="g")
        self.assertEqual(self.client.get("/logistica/garin").status_code, 200)
        self.assertEqual(self.client.get("/logistica").status_code, 403)
        self.assertEqual(self.client.get("/logistica/stock").status_code, 403)

    def test_selector_principal_muestra_portal_logistica(self):
        page = self.client.get("/")
        self.assertEqual(page.status_code, 200)
        self.assertIn("Logística".encode("utf-8"), page.data)
        self.assertIn(b'href="/logistica/login"', page.data)

    def test_login_local_no_habilita_dabra_y_permanece_oculto_por_defecto(self):
        page = self.client.get("/logistica/login")
        self.assertNotIn(b'name="password"', page.data)
        with self.client.session_transaction() as sess:
            sess["_csrf_token"] = "csrf-test"
        response = self.post("/logistica/login", {"usuario": "dabra", "password": "test-dabra"})
        self.assertEqual(response.status_code, 403)
        tecman.app.config["LOGISTICA_LOCAL_LOGIN_ENABLED"] = True
        response = self.post("/logistica/login", {"usuario": "dabra", "password": "test-dabra"})
        self.assertEqual(response.status_code, 200)
        with self.client.session_transaction() as sess:
            self.assertNotIn("logistica_role", sess)
        response = self.post("/logistica/login", {"usuario": "garin", "password": "test-garin"})
        self.assertTrue(response.headers["Location"].endswith("/logistica/garin"))

    def test_dabra_exige_sesion_entra_del_preparador_exacto(self):
        with self.client.session_transaction() as sess:
            sess.clear(); sess.update(logistica_role="dabra", logistica_user="esoria@grupodexter.com.ar", auth_provider="entra")
        self.assertEqual(self.client.get("/logistica").status_code, 403)
        with self.client.session_transaction() as sess:
            sess.clear(); sess.update(logistica_role="dabra", logistica_user="soria_demo", auth_provider="local_logistica")
        self.assertEqual(self.client.get("/logistica/stock").status_code, 403)
        self.session(logistica_role="dabra")
        self.assertEqual(self.client.get("/logistica").status_code, 200)
        self.assertEqual(self.client.get("/logistica/stock").status_code, 200)
        self.assertEqual(self.client.get("/logistica/garin").status_code, 403)

    def test_allowlist_dabra_legacy_no_revive_soria_y_grupos_existentes_siguen(self):
        with patch.dict(os.environ, {
            "LOGISTICA_ENTRA_DABRA_EMAILS": "esoria@grupodexter.com.ar,otro@grupodexter.com.ar",
            "LOGISTICA_ENTRA_DABRA_GROUP_ID": "grupo-dabra-viejo",
            "LOGISTICA_ENTRA_GARIN_GROUP_ID": "grupo-garin",
        }, clear=False):
            self.assertEqual(logistics_entra_role({"email": tecman.PREPARADOR_DABRA_EMAIL}), "dabra")
            self.assertIsNone(logistics_entra_role({"email": "esoria@grupodexter.com.ar", "claims": {"groups": ["grupo-dabra-viejo"]}}))
            self.assertIsNone(logistics_entra_role({"email": "otro@grupodexter.com.ar", "claims": {"groups": ["grupo-dabra-viejo"]}}))
            self.assertEqual(logistics_entra_role({"email": "garin@grupodexter.com.ar", "claims": {"groups": ["grupo-garin"]}}), "garin")

    def test_callback_entra_crea_sesion_dabra_solo_para_hdiosque(self):
        identity = {
            "object_id": "entra-hdiosque",
            "tenant_id": "tenant-test",
            "email": tecman.PREPARADOR_DABRA_EMAIL,
            "name": "Héctor Diosque",
            "claims": {},
        }
        msal_app = Mock()
        msal_app.acquire_token_by_authorization_code.return_value = {
            "id_token": "token-test",
            "access_token": "access-test",
        }
        with self.client.session_transaction() as sess:
            sess["entra_state"] = "state-test"
            sess["entra_nonce"] = "nonce-test"
            sess["entra_requested_portal"] = "logistica"
        with patch.object(tecman, "_entra_is_configured", return_value=True), \
             patch.object(tecman, "_create_msal_app", return_value=msal_app), \
             patch.object(tecman, "_validate_entra_id_token", return_value=identity), \
             patch.object(tecman, "_entra_group_ids", return_value=set()), \
             patch.object(tecman, "_audit_event"):
            response = self.client.get("/auth/entra/callback?state=state-test&code=code-test")
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.headers["Location"].endswith("/logistica"))
        with self.client.session_transaction() as sess:
            self.assertEqual(sess["auth_provider"], "entra")
            self.assertEqual(sess["logistica_user"], tecman.PREPARADOR_DABRA_EMAIL)
            self.assertEqual(sess["logistica_name"], tecman.PREPARADOR_DABRA_NAME)
            self.assertEqual(sess["logistica_role"], "dabra")

    def test_feature_flag_sucursal_apagada_bloquea_get_y_post(self):
        self.session(suc_user="suc011", suc_nombre="Sucursal 011")
        self.assertEqual(self.client.get("/logistica/sucursal").status_code, 404)
        self.assertEqual(self.post("/logistica/sucursal", {"item": "Guantes", "cantidad": "1"}).status_code, 404)
        self.assertEqual(self.post("/logistica/sucursal/pedidos/x/recepcion").status_code, 404)

    def test_sucursal_activada_requiere_sesion_valida_y_csrf_para_crear(self):
        tecman.app.config["LOGISTICA_INSUMOS_SUCURSALES_ENABLED"] = True
        self.assertEqual(self.client.get("/logistica/sucursal").status_code, 403)
        self.session(suc_user="suc011", suc_nombre="Sucursal 011")
        self.assertEqual(self.client.post("/logistica/sucursal", data={"item": "Guantes", "cantidad": "2"}).status_code, 400)
        response = self.post("/logistica/sucursal", {"item": "Guantes", "cantidad": "2"})
        self.assertEqual(response.status_code, 302)
        orders = list(tecman.logistica_service.state()["orders"].values())
        self.assertEqual([(o["sucursal"], o["lines"][0]["requested"]) for o in orders], [("Sucursal 011", 2)])

    def test_convalidacion_crea_requisicion_y_envia_email_mock_una_vez(self):
        service = tecman.logistica_service
        service.import_stock(b"item,cantidad\nGuantes,1\n", "s.csv", "Dabra")
        order = service.create_order("Sucursal 011", [{"item": "Guantes", "requested": 4}], "Sucursal")
        self.session(logistica_role="dabra", logistica_user="d", logistica_name="Dabra")
        self.assertEqual(self.post(f"/logistica/pedidos/{order['id']}", {"approved_0": "4"}).status_code, 302)
        self.assertEqual(self.post(f"/logistica/pedidos/{order['id']}", {"approved_0": "4"}).status_code, 302)
        state = service.state()
        requisition = next(iter(state["requisitions"].values()))
        self.assertEqual(requisition["email_status"], "sent")
        self.assertEqual(requisition["email_attempts"], 1)
        self.assertEqual(requisition["number"], "REQ-000001")
        self.smtp_mock.assert_called_once()
        recipients, subject, body = self.smtp_mock.call_args.args
        self.assertTrue(recipients)
        self.assertIn("REQ-000001", subject)
        self.assertIn("Guantes", body)

    def test_falla_smtp_ruta_conserva_requisicion_y_reintenta_con_csrf(self):
        service = tecman.logistica_service
        order = service.create_order("Sucursal 011", [{"item": "Barbijos", "requested": 2}], "Sucursal")
        self.smtp_mock.side_effect = RuntimeError("SMTP mock caído")
        self.session(logistica_role="dabra", logistica_user="d", logistica_name="Dabra")
        self.post(f"/logistica/pedidos/{order['id']}", {"approved_0": "2"})
        requisition = next(iter(service.state()["requisitions"].values()))
        self.assertEqual(requisition["email_status"], "failed")
        retry_url = f"/logistica/requisiciones/{requisition['id']}/reintentar-email"
        self.assertEqual(self.client.post(retry_url, data={}).status_code, 400)
        self.smtp_mock.side_effect = None
        self.assertEqual(self.post(retry_url).status_code, 302)
        self.assertEqual(service.state()["requisitions"][requisition["id"]]["email_status"], "sent")
        self.assertEqual(self.smtp_mock.call_count, 2)

    def test_garin_no_cambia_cantidades_ni_descontar_stock(self):
        service = tecman.logistica_service
        service.import_stock(b"item,cantidad\nGuantes,5\n", "s.csv", "Dabra")
        order = service.create_order("Sucursal 011", [{"item": "Guantes", "requested": 4}], "Test")
        service.approve(order["id"], {"Guantes": 4}, "Dabra")
        wave = service.create_wave([order["id"]], "semanal", "Dabra")
        self.session(logistica_role="garin", logistica_user="g", logistica_name="Garín")
        data = {"status": "retirado", "route": "R1", "observation": "ok", "approved_0": "999", "cantidad": "999"}
        self.assertEqual(self.post(f"/logistica/garin/olas/{wave['id']}", data).status_code, 302)
        state = service.state()
        self.assertEqual(state["orders"][order["id"]]["lines"][0]["approved"], 4)
        self.assertEqual(state["stock"]["Guantes"], 5)
        self.assertEqual(state["waves"][wave["id"]]["status"], "preparado")

        self.session(logistica_role="dabra", logistica_user="d", logistica_name="Dabra")
        self.assertEqual(self.post(f"/logistica/olas/{wave['id']}/entregar-garin").status_code, 302)
        self.assertEqual(self.post(f"/logistica/olas/{wave['id']}/entregar-garin").status_code, 302)
        state = service.state()
        self.assertEqual(state["stock"]["Guantes"], 1)
        self.assertEqual(state["waves"][wave["id"]]["status"], "retirado")
        self.assertEqual(len(state["movements"]), 1)

    def test_dabra_puede_crear_pedido_simulado_sin_habilitar_sucursales(self):
        self.session(logistica_role="garin", logistica_user="g")
        self.assertEqual(self.post("/logistica/demo/pedidos", {
            "sucursal": "Sucursal TEST", "item": "Guantes", "cantidad": "2"
        }).status_code, 403)
        self.session(logistica_role="dabra", logistica_user="d")
        response = self.post("/logistica/demo/pedidos", {
            "sucursal": "Sucursal TEST", "item": "Guantes", "cantidad": "2"
        })
        self.assertEqual(response.status_code, 302)
        orders = list(tecman.logistica_service.state()["orders"].values())
        self.assertEqual([(o["sucursal"], o["lines"][0]["requested"]) for o in orders], [("Sucursal TEST", 2)])
        self.assertFalse(tecman.app.config["LOGISTICA_INSUMOS_SUCURSALES_ENABLED"])
        tecman.app.config["LOGISTICA_PORTAL_TEST_MODE"] = False
        self.assertEqual(self.post("/logistica/demo/pedidos", {
            "sucursal": "Sucursal 014", "item": "Barbijos", "cantidad": "3"
        }).status_code, 404)
        self.assertEqual(len(tecman.logistica_service.state()["orders"]), 1)

    def test_sucursal_activada_solo_ve_y_confirma_sus_pedidos(self):
        service = tecman.logistica_service
        service.import_stock(b"item,cantidad\nGuantes,10\n", "s.csv", "Dabra")
        own = service.create_order("Sucursal 011", [{"item": "Guantes", "requested": 3}], "Test")
        other = service.create_order("Sucursal 014", [{"item": "Guantes", "requested": 2}], "Test")
        for order in (own, other):
            service.approve(order["id"], {"Guantes": order["lines"][0]["requested"]}, "Dabra")
            wave = service.create_wave([order["id"]], "semanal", "Dabra")
            for status in ("retirado", "en_distribucion", "entregado"):
                service.update_wave(wave["id"], status, "R", "", "Garín")
        tecman.app.config["LOGISTICA_INSUMOS_SUCURSALES_ENABLED"] = True
        self.session(suc_user="suc011", suc_nombre="Sucursal 011")
        page = self.client.get("/logistica/sucursal")
        self.assertEqual(page.status_code, 200)
        self.assertIn(own["id"].encode(), page.data)
        self.assertNotIn(other["id"].encode(), page.data)
        self.assertEqual(self.post(f"/logistica/sucursal/pedidos/{other['id']}/recepcion").status_code, 403)
        self.assertEqual(self.post(f"/logistica/sucursal/pedidos/{own['id']}/recepcion").status_code, 302)
        self.assertEqual(service.state()["orders"][own["id"]]["status"], "recibido")


if __name__ == "__main__":
    unittest.main()
