import copy
import hashlib
import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook


_TEST_ROOT = tempfile.TemporaryDirectory()
os.environ["TECMAN_DATA_DIR"] = str(Path(_TEST_ROOT.name) / "data")
os.environ["TECMAN_UPLOADS_DIR"] = str(Path(_TEST_ROOT.name) / "uploads")
os.environ.pop("DATABASE_URL", None)

import app as tecman  # noqa: E402


class CotizacionReparacionesTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        tecman.app.config.update(TESTING=True, SECRET_KEY="test-only", SERVER_NAME="localhost")
        tecman.USE_DB = False
        tecman.TICKETS_FILE = root / "tickets.json"
        tecman.NOTIF_ADMIN_FILE = root / "notif_admin.json"
        tecman.UPLOADS_DIR = root / "uploads"
        tecman.UPLOADS_DIR.mkdir(parents=True)
        tecman.save_notif_admin({"notificaciones": [{"id": "keep", "leida": False}]})
        tecman.save_tickets(self._tickets())
        self.client = tecman.app.test_client()
        self._admin_session()

    def tearDown(self):
        tecman.USE_DB = False
        self.temp.cleanup()

    @staticmethod
    def _tickets():
        base = {
            "categoria": "Reparaciones",
            "subcategoria": "General",
            "descripcion": "Reparar persiana",
            "prioridad": 2,
            "estado": "Nuevo",
            "creado": "2026-09-01T09:30:00",
            "actualizado": "2026-09-02T10:00:00",
            "observaciones": "Coordinar antes de visitar",
            "fotos": [],
        }
        return [
            dict(base, id=101, sucursal="Sucursal 120", asignado="Julio Fuga (JRF)", fotos=["foto 1.jpg"], notas=[{"autor": "Admin", "fecha": "2026-09-02T10:00:00", "texto": "Revisar motor"}]),
            dict(base, id=102, sucursal="Sucursal 120", asignado="Carolina", proveedor_nombre="JRF", zona_afectada="Salón"),
            dict(base, id=103, sucursal="Sucursal 126", asignado="Carolina", asignado_proveedor="fuga", descripcion="=2+2 no debe ser fórmula", presupuestos=[{"archivo": "presupuesto.pdf"}]),
            dict(base, id=104, sucursal="Sucursal 126", asignado="Ismael Allende (JRF)"),
            dict(base, id=105, sucursal="Sucursal 139", asignado="Julio Fuga", estado="Resuelto"),
            dict(base, id=106, sucursal="Sucursal 139", proveedor_nombre="Ismael Allende", estado="Cerrado"),
            dict(base, id=107, sucursal="Sucursal 145", asignado_proveedor="ismael", estado="Rechazado"),
            dict(base, id=201, sucursal="Sucursal 120", asignado="Otro proveedor"),
            dict(base, id=202, sucursal="Sucursal 120", proveedor_nombre="JRF Servicios Integrales"),
            dict(base, id=203, sucursal="Sucursal 120", asignado="No Julio Fuga"),
            dict(base, id=301, sucursal="Sucursal 120", asignado="Carolina", proveedor_nombre="CEYH", categoria="Electricidad", subcategoria="Iluminación", prioridad=1, creado="2026-09-10T08:00:00", descripcion="Cambiar reflector del depósito", zona_afectada="Depósito", observaciones="Acceso por portón lateral"),
            dict(base, id=302, sucursal="Sucursal 126", asignado="CEYH", categoria="Electricidad", subcategoria="Tablero", estado="Cerrado", creado="2026-08-10T08:00:00", fotos=["tablero.jpg"]),
            dict(base, id=401, sucursal="Sucursal 130", asignado="", proveedor_nombre="", asignado_proveedor="", responsable="Agustín Brahim", descripcion="Filtración sin asignar"),
            dict(base, id=501, sucursal="Sucursal 203", asignado="Gustavo Avellaneda", categoria="Sanitarios", subcategoria="Pérdida"),
            dict(base, id=601, sucursal="Sucursal 204", asignado="Carolina", proveedor_presupuesto="Presu SRL", categoria="Presupuestos"),
        ]

    def _admin_session(self, csrf="csrf-test"):
        with self.client.session_transaction() as session:
            session.clear()
            session["user"] = "admin-test"
            session["rol"] = "admin"
            session["nombre"] = "Agustín"
            session["_csrf_token"] = csrf

    @staticmethod
    def _digest(path):
        return hashlib.sha256(Path(path).read_bytes()).hexdigest()

    @staticmethod
    def _default_filters():
        values = {key: "" for key in tecman.QUOTE_FILTER_KEYS}
        values.update(proveedor_actual="Julio Fuga (JRF)", estado_scope="abiertos", adjuntos="todos")
        return values

    def _post(self, ids, **extra):
        data = {
            "_csrf_token": "csrf-test",
            "proveedor_destino": "Nuevo Proveedor SRL",
            "ticket_ids": [str(value) for value in ids],
            **self._default_filters(),
        }
        data.update(extra)
        return self.client.post("/admin/cotizacion-reparaciones", data=data)

    def test_default_jrf_usa_aliases_exactos_y_excluye_finalizados(self):
        filters = tecman._quote_filter_values({}, apply_initial_defaults=True)
        open_ids = [ticket["id"] for ticket in tecman._quote_filter_tickets(self._tickets(), filters)]
        filters["estado_scope"] = "todos"
        all_ids = [ticket["id"] for ticket in tecman._quote_filter_tickets(self._tickets(), filters)]

        self.assertEqual(open_ids, [101, 102, 103, 104])
        self.assertEqual(all_ids, [101, 102, 103, 104, 105, 106, 107])
        self.assertNotIn(202, all_ids)
        self.assertNotIn(203, all_ids)
        self.assertEqual(tecman.TICKET_FINAL_STATES, {"Rechazado", "Resuelto", "Cerrado"})

    def test_render_inicial_general_y_limpiar_filtros(self):
        response = self.client.get("/admin/cotizacion-reparaciones")
        self.assertEqual(response.status_code, 200)
        page = response.get_data(as_text=True)

        self.assertIn("Filtrá tickets de cualquier proveedor", page)
        self.assertIn("4 tickets", page)
        self.assertEqual(page.count('<section class="card quote-branch" data-branch-group>'), 2)
        self.assertEqual(page.count('type="checkbox" data-select-branch'), 2)
        for ticket_id in (101, 102, 103, 104):
            self.assertIn(f'value="{ticket_id}"', page)
        for ticket_id in (105, 301, 401, 501):
            self.assertNotIn(f'value="{ticket_id}"', page)
        self.assertIn('<option value="Julio Fuga (JRF)" selected>', page)
        self.assertIn('name="_csrf_token" value="csrf-test"', page)
        self.assertIn("Limpiar filtros", page)

        cleared = self.client.get("/admin/cotizacion-reparaciones?proveedor_actual=&estado_scope=todos&adjuntos=todos").get_data(as_text=True)
        self.assertIn("15 tickets", cleared)
        for ticket_id in (105, 301, 302, 401, 501, 601):
            self.assertIn(f'value="{ticket_id}"', cleared)

    def test_otro_proveedor_sin_proveedor_y_catalogo_completo(self):
        ceyh = self.client.get("/admin/cotizacion-reparaciones?proveedor_actual=CEYH&estado_scope=todos&adjuntos=todos").get_data(as_text=True)
        self.assertIn('value="301"', ceyh)
        self.assertIn('value="302"', ceyh)
        self.assertNotIn('value="101"', ceyh)

        unassigned = self.client.get(f"/admin/cotizacion-reparaciones?proveedor_actual={tecman.QUOTE_NO_PROVIDER}&estado_scope=todos&adjuntos=todos").get_data(as_text=True)
        self.assertIn('value="401"', unassigned)
        self.assertNotIn('value="301"', unassigned)

        self.assertIn('<option value="Gustavo Avellaneda"', ceyh)
        self.assertIn('<option value="Martin Microglobal"', ceyh)
        self.assertIn('<option value="Otro proveedor"', ceyh)
        self.assertIn('<option value="Presu SRL"', ceyh)
        self.assertIn('>Sin proveedor</option>', ceyh)

        presupuesto = self.client.get("/admin/cotizacion-reparaciones?proveedor_actual=Presu+SRL&estado_scope=abiertos&adjuntos=todos").get_data(as_text=True)
        self.assertIn('value="601"', presupuesto)
        self.assertNotIn('value="301"', presupuesto)

    def test_filtros_combinados_server_side(self):
        filters = tecman._quote_filter_values({
            "proveedor_actual": "CEYH",
            "sucursal": "Sucursal 120",
            "estado_scope": "abiertos",
            "categoria": "Electricidad",
            "subcategoria": "Iluminación",
            "prioridad": "1",
            "fecha_desde": "2026-09-05",
            "fecha_hasta": "2026-09-15",
            "antiguedad_min": "8",
            "antiguedad_max": "8",
            "responsable": "Carolina",
            "adjuntos": "sin",
            "q": "deposito porton",
        })
        result = tecman._quote_filter_tickets(
            self._tickets(), filters, now=tecman.datetime.datetime(2026, 9, 18, 10, 0)
        )
        self.assertEqual([ticket["id"] for ticket in result], [301])

        filters["adjuntos"] = "con"
        self.assertEqual(tecman._quote_filter_tickets(self._tickets(), filters, now=tecman.datetime.datetime(2026, 9, 18, 10, 0)), [])

    def test_busqueda_y_filtros_de_adjunto(self):
        search = self.client.get("/admin/cotizacion-reparaciones?proveedor_actual=&estado_scope=todos&adjuntos=todos&q=porton+lateral").get_data(as_text=True)
        self.assertIn('value="301"', search)
        self.assertNotIn('value="302"', search)

        with_files = self.client.get("/admin/cotizacion-reparaciones?proveedor_actual=CEYH&estado_scope=todos&adjuntos=con").get_data(as_text=True)
        self.assertIn('value="302"', with_files)
        self.assertNotIn('value="301"', with_files)

        note_search = self.client.get("/admin/cotizacion-reparaciones?proveedor_actual=Julio+Fuga+(JRF)&estado_scope=abiertos&adjuntos=todos&q=motor").get_data(as_text=True)
        self.assertIn('value="101"', note_search)
        self.assertNotIn('value="102"', note_search)

    def test_xlsx_incluye_origen_filtros_destino_columnas_y_enlaces(self):
        before_tickets = self._digest(tecman.TICKETS_FILE)
        before_notices = self._digest(tecman.NOTIF_ADMIN_FILE)
        response = self._post([101, 103])

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        workbook = load_workbook(io.BytesIO(response.data), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Resumen", "Sucursal 120", "Sucursal 126"])

        summary = workbook["Resumen"]
        self.assertEqual(summary["A1"].value, "Cotización de reparaciones")
        self.assertEqual(summary["B2"].value, "Julio Fuga (JRF)")
        self.assertEqual(summary["B3"].value, "Nuevo Proveedor SRL")
        self.assertEqual(summary["B5"].value, 2)
        self.assertIn("Estado: Abiertos", summary["B6"].value)
        self.assertEqual(summary["A9"].value, "Sucursal 120")
        self.assertEqual(summary["B9"].value, 1)

        sheet_120 = workbook["Sucursal 120"]
        headers = [cell.value for cell in sheet_120[7]]
        self.assertEqual(headers[:14], [
            "Sucursal", "Ticket", "Proveedor actual", "Responsable / asignado", "Categoría",
            "Subcategoría", "Descripción / problema", "Zona / ubicación", "Prioridad", "Estado",
            "Fecha", "Antigüedad (días)", "Observaciones relevantes", "Ver ticket",
        ])
        self.assertIn("Adjunto 1", headers)
        self.assertEqual(sheet_120.max_row, 8)
        self.assertEqual(sheet_120["B8"].value, "101")
        self.assertEqual(sheet_120["C8"].value, "Julio Fuga (JRF)")
        self.assertEqual(sheet_120["B2"].value, "Julio Fuga (JRF)")
        self.assertEqual(sheet_120["B3"].value, "Nuevo Proveedor SRL")
        self.assertEqual(sheet_120["N8"].hyperlink.target, "http://localhost/admin/ticket/101")
        self.assertEqual(sheet_120["O8"].hyperlink.target, "http://localhost/static/uploads/foto%201.jpg")
        self.assertIn("Revisar motor", sheet_120["M8"].value)

        sheet_126 = workbook["Sucursal 126"]
        self.assertEqual(sheet_126["B8"].value, "103")
        self.assertEqual(sheet_126["G8"].value, "'=2+2 no debe ser fórmula")
        self.assertEqual(sheet_126["O8"].hyperlink.target, "http://localhost/static/uploads/presupuesto.pdf")
        workbook.close()

        self.assertEqual(self._digest(tecman.TICKETS_FILE), before_tickets)
        self.assertEqual(self._digest(tecman.NOTIF_ADMIN_FILE), before_notices)
        self.assertEqual(list(tecman.UPLOADS_DIR.iterdir()), [])

    def test_post_reconstruye_filtros_y_rechaza_ids_ajenos_invalidos(self):
        cases = [
            ([], {}, 400),
            ([301], {}, 400),
            ([105], {}, 400),
            ([101, 101], {}, 400),
            ([101], {"_csrf_token": "incorrecto"}, 400),
            ([101], {"proveedor_destino": "x" * 121}, 400),
            ([101], {"fecha_desde": "18/09/2026"}, 400),
            ([101], {"antiguedad_min": "20", "antiguedad_max": "10"}, 400),
        ]
        for ids, extra, expected in cases:
            with self.subTest(ids=ids, extra=extra):
                before_tickets = self._digest(tecman.TICKETS_FILE)
                before_notices = self._digest(tecman.NOTIF_ADMIN_FILE)
                response = self._post(ids, **extra)
                self.assertEqual(response.status_code, expected)
                self.assertEqual(self._digest(tecman.TICKETS_FILE), before_tickets)
                self.assertEqual(self._digest(tecman.NOTIF_ADMIN_FILE), before_notices)

        valid_other = self._post([301], proveedor_actual="CEYH")
        self.assertEqual(valid_other.status_code, 200)
        valid_closed = self._post([105], estado_scope="cerrados")
        self.assertEqual(valid_closed.status_code, 200)

    def test_sesion_db_y_nombres_de_hoja(self):
        with self.client.session_transaction() as session:
            session.clear()
        no_session = self.client.get("/admin/cotizacion-reparaciones")
        self.assertEqual(no_session.status_code, 302)
        self.assertIn("/admin/login", no_session.headers["Location"])
        self._admin_session()

        tickets = copy.deepcopy(self._tickets())
        ticket_model = object()
        with patch.object(tecman, "TicketDB", ticket_model, create=True), \
             patch.object(tecman, "_db_list", return_value=copy.deepcopy(tickets)) as db_list, \
             patch.object(tecman, "_db_replace") as db_replace, \
             patch.object(tecman, "USE_DB", True):
            response = self.client.get("/admin/cotizacion-reparaciones")
        self.assertEqual(response.status_code, 200)
        db_list.assert_called_once_with(ticket_model)
        db_replace.assert_not_called()

        weird = [
            dict(tickets[0], id=601, sucursal="Sucursal [Norte]/Nombre extremadamente largo 1234567890"),
            dict(tickets[1], id=602, sucursal="Sucursal :Norte?/Nombre extremadamente largo 1234567890"),
        ]
        with tecman.app.test_request_context("/", base_url="http://localhost"):
            workbook = tecman._build_quote_workbook(weird, "", generated_at=tecman.datetime.datetime(2026, 9, 18, 10, 0))
        self.assertEqual(len(workbook.sheetnames), len(set(name.casefold() for name in workbook.sheetnames)))
        for name in workbook.sheetnames:
            self.assertLessEqual(len(name), 31)
            self.assertFalse(any(char in name for char in "[]:*?/\\"))
        workbook.close()


if __name__ == "__main__":
    unittest.main()
