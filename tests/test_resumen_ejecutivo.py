import copy
import hashlib
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

_TEST_ROOT = tempfile.TemporaryDirectory()
os.environ["TECMAN_DATA_DIR"] = str(Path(_TEST_ROOT.name) / "data")
os.environ["TECMAN_UPLOADS_DIR"] = str(Path(_TEST_ROOT.name) / "uploads")
os.environ.pop("DATABASE_URL", None)

import app as tecman  # noqa: E402
from executive_report import build_report, build_workbook, parse_period  # noqa: E402


class ResumenEjecutivoTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        tecman.app.config.update(TESTING=True, SECRET_KEY="test", SERVER_NAME="localhost")
        tecman.USE_DB = False
        tecman.TICKETS_FILE = root / "tickets.json"
        tecman.AUTH_AUDIT_FILE = root / "auth_audit_events.json"
        tecman.USERS_FILE = root / "users.json"
        tecman.UPLOADS_DIR = root / "uploads"
        tecman.UPLOADS_DIR.mkdir(parents=True)
        tecman.TICKETS_FILE.write_text(json.dumps(self.tickets()), encoding="utf-8")
        tecman.AUTH_AUDIT_FILE.write_text(json.dumps({"events": self.events()}), encoding="utf-8")
        tecman.USERS_FILE.write_text(json.dumps({"users": self.users()}), encoding="utf-8")
        self.client = tecman.app.test_client()

    def tearDown(self):
        tecman.USE_DB = False
        self.temp.cleanup()

    @staticmethod
    def catalog():
        return {
            "001": {"tienda": "Sucursal =Uno"},
            "002": {"tienda": "Sucursal Dos"},
            "003": {"tienda": "Sucursal Tres"},
        }

    @staticmethod
    def tickets():
        base = {"categoria": "Mantenimiento", "subcategoria": "General", "tipo": "Correctivo", "estado": "Nuevo", "asignado": "Equipo", "sector": "Operaciones"}
        return [
            dict(base, id="seed", sucursal="Sucursal 001", creado="2025-01-01T09:00:00", seed=True),
            dict(base, id="tutorial", sucursal="Sucursal 001", creado="2025-01-02T09:00:00", origen="tutorial"),
            dict(base, id="demo", sucursal="Sucursal 001", creado="2025-01-03T09:00:00", descripcion="DEMO guiada"),
            dict(base, id="prueba", sucursal="Sucursal 001", creado="2025-01-04T09:00:00", descripcion="TICKET PRUEBA - exacto"),
            dict(base, id=1, sucursal="Sucursal 001", creado="2026-01-01T08:00:00", actualizado="2026-01-01T10:00:00"),
            dict(base, id="legacy-2", sucursal_num="1", creado="2026-01-02T08:00:00", categoria="Materiales", tipo="insumos", estado="Cerrado", fecha_cierre="2026-01-03T08:00:00", notas=[{"fecha": "2026-01-02T09:00:00", "texto": "secreto"}, {"fecha": "2026-01-04T09:00:00", "tipo": "compra_np_derivado_compras"}], asignado="=RESP"),
            dict(base, id=3, sucursal="Sucursal 002", creado="2026-01-03T08:00:00", estado="Resuelto", fecha_cierre="2026-01-05T08:00:00", historial=[{"fecha": "2026-01-03T12:00:00", "detalle": "privado"}]),
            dict(base, id=4, sucursal="Sucursal 002", creado="2026-01-04T08:00:00", estado="Pendiente", asignado="=RESP", historial=[{"fecha": "2026-01-04T14:00:00"}], adjunto_obligatorio=True, adjuntos=[]),
            dict(base, id=5, sucursal="Sucursal 999", creado="2026-01-05T08:00:00", estado="Rechazado"),
            dict(base, id=6, sucursal="sin código", creado="2026-02-15T08:00:00", estado="Nuevo"),
        ]

    @staticmethod
    def events():
        return [
            {"event_type": "login_failed", "user_id": "u1", "created_at": "2025-12-01T00:00:00"},
            {"event_type": "login_success", "user_id": "u1", "created_at": "2025-12-31T00:00:00", "details": {}},
            {"event_type": "login_success", "user_id": "u1", "created_at": "2026-01-10T00:00:00", "details": {}},
            {"event_type": "login_success", "user_id": "u2", "created_at": "2026-02-01T00:00:00", "details": {"sucursal_num": "002"}},
        ]

    @staticmethod
    def users():
        return [{"id": "u1", "sucursal_num": "001", "email": "privado@example.com"}, {"id": "u2", "role": "sucursal"}]

    def report(self, **kwargs):
        return build_report(self.tickets(), self.catalog(), self.events(), self.users(), now=tecman.datetime.datetime(2026, 2, 10), **kwargs)

    def _session(self, role="admin"):
        with self.client.session_transaction() as session:
            session.clear(); session["user"] = "test"; session["rol"] = role; session["nombre"] = "Test"

    @staticmethod
    def digest(path):
        return hashlib.sha256(Path(path).read_bytes()).hexdigest()

    def test_periodo_automatico_exclusiones_totales_top_y_catalogo(self):
        report = self.report()
        self.assertEqual(report["periodo"], {"desde": "2026-01-01", "hasta_exclusivo": "2026-01-31", "fuente": "automatico", "dias": 30})
        self.assertEqual(report["metricas"]["total"], 5)
        self.assertEqual(report["metodologia"]["excluidos_total"], 4)
        self.assertEqual(report["sucursales_resumen"], {"catalogadas": 3, "con_tickets": 2, "sin_tickets": 1})
        self.assertEqual(report["top_sucursales"][0]["cantidad"], 2)
        self.assertTrue(all(report["reconciliaciones"].values()))
        self.assertLessEqual(report["metricas"]["concentracion_top5"], report["metricas"]["concentracion_top10"])

    def test_query_estricta_limite_y_semantica_hasta_exclusiva(self):
        with self.assertRaises(ValueError): parse_period({"desde": "2026-01-01"}, self.tickets())
        with self.assertRaises(ValueError): parse_period({"desde": "01/01/2026", "hasta": "2026-02-01"}, self.tickets())
        with self.assertRaises(ValueError): parse_period({"desde": "2026-01-01", "hasta": "2027-01-03"}, self.tickets())
        report = self.report(query={"desde": "2026-01-02", "hasta": "2026-01-05"})
        self.assertEqual(report["metricas"]["total"], 3)

    def test_estados_tiempos_backlog_compras_y_calidad(self):
        report = self.report()
        self.assertEqual(report["metricas"]["finalizados"], 2)
        self.assertEqual(report["metricas"]["compras_insumos"], 1)
        self.assertEqual(report["metricas"]["derivaciones_compras_insumos"], 1)
        self.assertEqual(report["metricas"]["backlog"], 2)
        self.assertEqual(report["tiempos"]["primera_gestion"]["estado"], "disponible")
        self.assertEqual(report["tiempos"]["resolucion"]["estado"], "sin datos suficientes")
        self.assertEqual(report["calidad"]["adjunto_obligatorio_faltante"], 1)
        self.assertEqual(report["calidad"]["sucursal_no_normalizada"], 1)
        self.assertNotIn("duplic", json.dumps(report).lower())
        self.assertEqual(len(report["recomendaciones"]), 8)
        self.assertEqual(len(report["metodologia"]["umbrales_recomendaciones"]), 8)

    def test_login_completo_parcial_y_ausente(self):
        complete = self.report()
        self.assertTrue(complete["login"]["cobertura_suficiente"])
        self.assertEqual(complete["login"]["etiqueta_ausencia"], "nunca ingresaron")
        partial = build_report(self.tickets(), self.catalog(), self.events()[1:2], self.users(), now=tecman.datetime.datetime(2026, 2, 10))
        self.assertEqual(partial["login"]["etiqueta_ausencia"], "sin evidencia de ingreso")
        absent = build_report(self.tickets(), self.catalog(), [], self.users(), now=tecman.datetime.datetime(2026, 2, 10))
        self.assertIsNone(absent["login"]["cobertura_desde"])
        self.assertEqual(absent["login"]["eventos_login_success_periodo"], 0)

    def test_login_mapea_email_canonico_sin_exponerlo(self):
        report = build_report(
            self.tickets(), self.catalog(),
            [{"event_type": "login_success", "user_id": "db-user", "created_at": "2026-01-10T00:00:00"}],
            [{"id": "db-user", "email": "sucursal1@example.com"}],
            branch_email_map={"sucursal1@example.com": "001"},
            now=tecman.datetime.datetime(2026, 2, 10),
        )
        self.assertEqual(report["login"]["sucursales_con_evidencia"], 1)
        self.assertNotIn("sucursal1@example.com", json.dumps(report))

    def test_cero_tickets_y_auditoria_ausente(self):
        report = build_report([], self.catalog(), [], [], query={})
        self.assertEqual(report["metricas"]["total"], 0)
        self.assertEqual(report["periodo"]["fuente"], "automatico_sin_datos")
        self.assertEqual(report["sucursales_resumen"]["sin_tickets"], 3)
        self.assertEqual(report["tiempos"]["resolucion"]["estado"], "sin datos suficientes")

    def test_xlsx_hojas_sanitizacion_y_sin_formulas(self):
        workbook = build_workbook(self.report())
        output = io.BytesIO(); workbook.save(output); workbook.close(); output.seek(0)
        reopened = load_workbook(output, data_only=False)
        self.assertEqual(reopened.sheetnames, ["Resumen", "Sucursales", "Categorias", "Estados", "Pendientes"])
        self.assertEqual(reopened["Sucursales"]["B2"].value, "Sucursal =Uno")
        self.assertTrue(any(cell.value == "'=RESP" for sheet in reopened.worksheets for row in sheet.iter_rows() for cell in row))
        for sheet in reopened.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")
                    if isinstance(cell.value, str):
                        self.assertNotIn("privado@example.com", cell.value)
                        self.assertNotIn("secreto", cell.value)
        reopened.close()

    def test_html_xlsx_auth_y_get_sin_mutacion(self):
        before = [self.digest(path) for path in (tecman.TICKETS_FILE, tecman.AUTH_AUDIT_FILE, tecman.USERS_FILE)]
        anonymous = self.client.get("/admin/resumen-ejecutivo")
        self.assertEqual(anonymous.status_code, 302)
        self._session("tecnico")
        self.assertEqual(self.client.get("/admin/resumen-ejecutivo").status_code, 403)
        self._session("admin")
        with patch("app._executive_report_sources", return_value=(self.tickets(), self.catalog(), self.events(), self.users())):
            html = self.client.get("/admin/resumen-ejecutivo")
            xlsx = self.client.get("/admin/resumen-ejecutivo.xlsx")
        self.assertEqual(html.status_code, 200)
        page = html.get_data(as_text=True)
        self.assertIn("Resumen ejecutivo", page); self.assertIn("2026-01-31", page); self.assertNotIn("privado@example.com", page); self.assertNotIn("secreto", page)
        self.assertEqual(xlsx.status_code, 200)
        self.assertEqual(xlsx.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertEqual(load_workbook(io.BytesIO(xlsx.data)).sheetnames, ["Resumen", "Sucursales", "Categorias", "Estados", "Pendientes"])
        after = [self.digest(path) for path in (tecman.TICKETS_FILE, tecman.AUTH_AUDIT_FILE, tecman.USERS_FILE)]
        self.assertEqual(before, after)
        self.assertEqual(list(tecman.UPLOADS_DIR.iterdir()), [])

    def test_json_y_db_simulado_usan_fuentes_sin_escribir(self):
        tickets, _, events, users = tecman._executive_report_sources()
        self.assertEqual(len(tickets), len(self.tickets())); self.assertEqual(len(events), len(self.events())); self.assertEqual(len(users), len(self.users()))

        class Query:
            def __init__(self, values): self.values = values
            def order_by(self, *args): return self
            def all(self): return self.values
        fake_event = SimpleNamespace(user_id="u1", event_type="login_success", details={"sucursal_num": "001"}, created_at=tecman.datetime.datetime(2026, 1, 2))
        fake_user = SimpleNamespace(id="u1", role="sucursal", sucursal_num="001", sucursal=None)
        column = SimpleNamespace(asc=lambda: None)
        fake_event_model = SimpleNamespace(query=Query([fake_event]), created_at=column)
        fake_user_model = SimpleNamespace(query=Query([fake_user]), created_at=column)
        with patch.object(tecman, "USE_DB", True), patch.object(tecman, "AuthAuditEventDB", fake_event_model, create=True), patch.object(tecman, "UserDB", fake_user_model, create=True), patch.object(tecman, "load_tickets", return_value=copy.deepcopy(self.tickets())) as loader:
            db_tickets, _, db_events, db_users = tecman._executive_report_sources()
        loader.assert_called_once_with(readonly=True)
        self.assertEqual(len(db_tickets), len(self.tickets())); self.assertEqual(db_events, [fake_event]); self.assertEqual(db_users, [fake_user])

    def test_load_tickets_readonly_no_normaliza_ni_persiste(self):
        raw = [{"id": "historico", "categoria": "Materiales", "asignado": "Proveedor viejo"}]
        before = copy.deepcopy(raw)
        with patch("app._load_tickets_raw", return_value=raw), patch("app.save_tickets") as save:
            loaded = tecman.load_tickets(readonly=True)
        self.assertEqual(loaded, before); save.assert_not_called()


if __name__ == "__main__":
    unittest.main()
