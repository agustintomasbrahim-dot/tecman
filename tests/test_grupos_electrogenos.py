import io
import os
import tempfile
import unittest
from unittest import mock
from werkzeug.datastructures import FileStorage

_TEST_DATA = tempfile.TemporaryDirectory()
os.environ["TECMAN_DATA_DIR"] = _TEST_DATA.name
os.environ["TECMAN_UPLOADS_DIR"] = os.path.join(_TEST_DATA.name, "uploads")
os.environ.pop("DATABASE_URL", None)

import app as tecman  # noqa: E402


class GruposElectrogenosTest(unittest.TestCase):
    def setUp(self):
        tecman.app.config.update(TESTING=True, SECRET_KEY="test-only")
        tecman.USE_DB = False
        tecman.GRUPOS_ELECTROGENOS_FILE = tecman.Path(_TEST_DATA.name) / "grupos_electrogenos.json"
        tecman.NOTIF_ADMIN_FILE = tecman.Path(_TEST_DATA.name) / "notif_admin.json"
        tecman.GRUPOS_ELECTROGENOS_UPLOADS_DIR = tecman.Path(_TEST_DATA.name) / "uploads" / "grupos_electrogenos"
        tecman.GRUPOS_ELECTROGENOS_UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
        for path in (tecman.GRUPOS_ELECTROGENOS_FILE, tecman.NOTIF_ADMIN_FILE):
            if path.exists():
                path.unlink()
        for path in tecman.GRUPOS_ELECTROGENOS_UPLOADS_DIR.iterdir():
            path.unlink()
        self.client = tecman.app.test_client()

    PNG_1X1 = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDAT\x08\xd7c\xf8\xcf\xc0\xf0\x1f\x00\x05\x00\x01\xff\x89\x99\r\x1d\x00\x00\x00\x00IEND\xaeB`\x82"
    )

    def _admin_session(self):
        with self.client.session_transaction() as sess:
            sess["user"] = "agustin"
            sess["nombre"] = "Admin Test"
            sess["rol"] = "admin"
            sess["_csrf_token"] = "csrf-test"

    def _sucursal_session(self, label):
        with self.client.session_transaction() as sess:
            sess.clear()
            sess["suc_user"] = "test"
            sess["suc_nombre"] = label
            sess["_csrf_token"] = "csrf-test"

    def _supervisor_session(self, scope=None):
        with self.client.session_transaction() as sess:
            sess.clear()
            sess["suc_user"] = "supervisor-test"
            sess["suc_nombre"] = "Portal Sucursales"
            sess["suc_general"] = True
            if scope is not None:
                sess["suc_scope_nums"] = scope
            sess["_csrf_token"] = "csrf-test"

    def _equipo(self, sucursal=None, marca="Equipo Test"):
        with tecman.app.test_request_context("/"):
            return tecman._generador_from_values(
                {"sucursal": sucursal or tecman.SUCURSALES[0], "marca": marca}, "Admin", "test"
            )

    def _post_novedad(self, equipo, tipo, **overrides):
        payload = {
            "_csrf_token": "csrf-test",
            "tipo": tipo,
            "fecha_evento": (tecman.datetime.datetime.now() - tecman.datetime.timedelta(minutes=1)).replace(second=0, microsecond=0).isoformat(timespec="minutes"),
            "observacion": "Detalle de prueba" if tipo == "problema_falla" else "",
            "proveedor_tecnico": "Técnico Test",
            "foto": (io.BytesIO(self.PNG_1X1), "equipo.png"),
        }
        payload.update(overrides)
        return self.client.post(
            f"/suc/grupos-electrogenos/{equipo['id']}/novedades",
            data=payload,
            content_type="multipart/form-data",
        )

    def test_admin_manual_csv_and_empty_template(self):
        self._admin_session()
        suc1, suc2 = tecman.SUCURSALES[:2]
        response = self.client.post("/admin/grupos-electrogenos", data={
            "_csrf_token": "csrf-test", "sucursal": suc1, "marca": "Marca A",
            "modelo": "Modelo 1", "numero_serie": "SER-1", "potencia": "50 kVA",
        })
        self.assertEqual(response.status_code, 302)

        csv_data = (
            "sucursal,marca,modelo,potencia,numero de serie,combustible,ubicacion,estado,ultima revision,proximo mantenimiento,proveedor,observaciones\n"
            f"{suc2},Marca B,Modelo 2,80 kVA,SER-2,Diesel,Patio,Operativo,2026-01-01,2026-10-01,Proveedor,\n"
        ).encode()
        response = self.client.post("/admin/grupos-electrogenos/importar", data={
            "_csrf_token": "csrf-test", "archivo": (io.BytesIO(csv_data), "equipos.csv")
        }, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 302)
        items = tecman.load_grupos_electrogenos()["grupos_electrogenos"]
        self.assertEqual(len(items), 2)
        self.assertTrue(all(x["estado_validacion"] == "pendiente_validacion" for x in items))
        self.assertTrue(all(x["historial"] for x in items))
        self.assertEqual(len(tecman.load_notif_admin()["notificaciones"]), 2)

        response = self.client.get("/admin/grupos-electrogenos/plantilla.xlsx")
        self.assertEqual(response.status_code, 200)
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(response.data), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(len(rows), 1)
        self.assertEqual(tuple(rows[0]), tecman.GENERADOR_IMPORT_HEADERS)

    def test_import_is_atomic_on_invalid_branch(self):
        self._admin_session()
        suc = tecman.SUCURSALES[0]
        csv_data = f"sucursal,marca\n{suc},OK\nSucursal inexistente,Error\n".encode()
        response = self.client.post("/admin/grupos-electrogenos/importar", data={
            "_csrf_token": "csrf-test", "archivo": (io.BytesIO(csv_data), "equipos.csv")
        }, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(tecman.load_grupos_electrogenos()["grupos_electrogenos"], [])

    def test_branch_isolation_and_audited_validation(self):
        suc1, suc2 = tecman.SUCURSALES[:2]
        with tecman.app.test_request_context("/"):
            one = tecman._generador_from_values({"sucursal": suc1, "marca": "Visible"}, "Admin", "test")
            two = tecman._generador_from_values({"sucursal": suc2, "marca": "Oculto"}, "Admin", "test")
        tecman.save_grupos_electrogenos({"grupos_electrogenos": [one, two]})
        self._sucursal_session(suc1)

        response = self.client.get("/suc/grupos-electrogenos")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Visible", response.data)
        self.assertNotIn(b"Oculto", response.data)

        forbidden = self.client.post(f"/suc/grupos-electrogenos/{two['id']}/validar", data={
            "_csrf_token": "csrf-test", "accion": "confirmar"
        })
        self.assertEqual(forbidden.status_code, 404)

        confirmed = self.client.post(f"/suc/grupos-electrogenos/{one['id']}/validar", data={
            "_csrf_token": "csrf-test", "accion": "confirmar"
        })
        self.assertEqual(confirmed.status_code, 302)
        saved = tecman.load_grupos_electrogenos()["grupos_electrogenos"]
        self.assertEqual(next(x for x in saved if x["id"] == one["id"])["estado_validacion"], "validado")

        accepted = self.client.post(f"/suc/grupos-electrogenos/{one['id']}/validar", data={
            "_csrf_token": "csrf-test", "accion": "diferencias", "diferencias": "Serie no coincide"
        })
        self.assertEqual(accepted.status_code, 302)
        saved = tecman.load_grupos_electrogenos()["grupos_electrogenos"]
        current = next(x for x in saved if x["id"] == one["id"])
        untouched = next(x for x in saved if x["id"] == two["id"])
        self.assertEqual(current["estado_validacion"], "con_diferencias")
        self.assertEqual(current["historial"][-1]["accion"], "diferencias_informadas")
        self.assertEqual(untouched["estado_validacion"], "pendiente_validacion")

    def test_xlsx_import(self):
        self._admin_session()
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(list(tecman.GENERADOR_IMPORT_HEADERS))
        sheet.append([tecman.SUCURSALES[0], "Marca XLSX", "Modelo XLSX", "30 kVA", "SER-XLSX", "Gas", "Exterior", "Reserva", "2026-01-01", "2027-01-01", "Proveedor", ""])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)
        response = self.client.post("/admin/grupos-electrogenos/importar", data={
            "_csrf_token": "csrf-test", "archivo": (output, "equipos.xlsx")
        }, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 302)
        items = tecman.load_grupos_electrogenos()["grupos_electrogenos"]
        self.assertEqual(items[0]["numero_serie"], "SER-XLSX")

    def test_general_session_without_scope_is_forbidden(self):
        suc = tecman.SUCURSALES[0]
        with tecman.app.test_request_context("/"):
            item = tecman._generador_from_values({"sucursal": suc, "marca": "No visible"}, "Admin", "test")
        tecman.save_grupos_electrogenos({"grupos_electrogenos": [item]})
        self._supervisor_session()
        self.assertEqual(self.client.get("/suc/grupos-electrogenos").status_code, 403)
        response = self.client.post(f"/suc/grupos-electrogenos/{item['id']}/validar", data={
            "_csrf_token": "csrf-test", "accion": "confirmar"
        })
        self.assertEqual(response.status_code, 403)

    def test_supervisor_with_scope_only_sees_and_validates_scope(self):
        suc1, suc2 = tecman.SUCURSALES[:2]
        with tecman.app.test_request_context("/"):
            one = tecman._generador_from_values({"sucursal": suc1, "marca": "En scope"}, "Admin", "test")
            two = tecman._generador_from_values({"sucursal": suc2, "marca": "Fuera scope"}, "Admin", "test")
        tecman.save_grupos_electrogenos({"grupos_electrogenos": [one, two]})
        self._supervisor_session([one["sucursal_num"]])
        response = self.client.get("/suc/grupos-electrogenos")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"En scope", response.data)
        self.assertNotIn(b"Fuera scope", response.data)
        forbidden = self.client.post(f"/suc/grupos-electrogenos/{two['id']}/validar", data={
            "_csrf_token": "csrf-test", "accion": "confirmar"
        })
        self.assertEqual(forbidden.status_code, 404)
        allowed = self.client.post(f"/suc/grupos-electrogenos/{one['id']}/validar", data={
            "_csrf_token": "csrf-test", "accion": "confirmar"
        })
        self.assertEqual(allowed.status_code, 302)

    def test_invalid_dates_are_rejected_atomically(self):
        self._admin_session()
        suc = tecman.SUCURSALES[0]
        manual = self.client.post("/admin/grupos-electrogenos", data={
            "_csrf_token": "csrf-test", "sucursal": suc, "ultima_revision": "15/09/2026"
        })
        self.assertEqual(manual.status_code, 302)
        self.assertEqual(tecman.load_grupos_electrogenos()["grupos_electrogenos"], [])

        csv_data = (
            "sucursal,marca,ultima revision,proximo mantenimiento\n"
            f"{suc},Valido,2026-01-01,2027-01-01\n"
            f"{suc},Invalido,2026-99-01,2027-01-01\n"
        ).encode()
        imported = self.client.post("/admin/grupos-electrogenos/importar", data={
            "_csrf_token": "csrf-test", "archivo": (io.BytesIO(csv_data), "equipos.csv")
        }, content_type="multipart/form-data")
        self.assertEqual(imported.status_code, 302)
        self.assertEqual(tecman.load_grupos_electrogenos()["grupos_electrogenos"], [])

    def test_manual_duplicate_serial_is_rejected_per_branch(self):
        self._admin_session()
        suc1, suc2 = tecman.SUCURSALES[:2]
        payload = {"_csrf_token": "csrf-test", "sucursal": suc1, "numero_serie": " DUP-01 "}
        self.assertEqual(self.client.post("/admin/grupos-electrogenos", data=payload).status_code, 302)
        payload["numero_serie"] = "dup-01"
        self.assertEqual(self.client.post("/admin/grupos-electrogenos", data=payload).status_code, 302)
        items = tecman.load_grupos_electrogenos()["grupos_electrogenos"]
        self.assertEqual(len(items), 1)
        payload.update({"sucursal": suc2, "numero_serie": "DUP-01"})
        self.assertEqual(self.client.post("/admin/grupos-electrogenos", data=payload).status_code, 302)
        self.assertEqual(len(tecman.load_grupos_electrogenos()["grupos_electrogenos"]), 2)

    def test_routes_are_protected(self):
        self.assertEqual(self.client.get("/admin/grupos-electrogenos").status_code, 302)
        self.assertEqual(self.client.get("/suc/grupos-electrogenos").status_code, 302)

    def test_post_requires_csrf(self):
        self._admin_session()
        response = self.client.post("/admin/grupos-electrogenos", data={"sucursal": tecman.SUCURSALES[0]})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(tecman.load_grupos_electrogenos()["grupos_electrogenos"], [])

    def test_inventory_2026_09_is_idempotent_and_preserves_existing_values(self):
        with tecman.app.test_request_context("/"):
            existing = tecman._generador_from_values({
                "sucursal": "Sucursal 077",
                "marca": "Marca confirmada por sucursal",
            }, "Admin", "test")
        tecman.save_grupos_electrogenos({"grupos_electrogenos": [existing]})

        with tecman.app.app_context():
            first = tecman._ensure_grupos_electrogenos_inventario_2026_09()
            second = tecman._ensure_grupos_electrogenos_inventario_2026_09()

        self.assertEqual(first, 21)
        self.assertEqual(second, 0)
        items = tecman.load_grupos_electrogenos()["grupos_electrogenos"]
        self.assertEqual(len(items), 21)
        by_branch = {item["sucursal_num"]: item for item in items}
        self.assertEqual(by_branch["077"]["marca"], "Marca confirmada por sucursal")
        self.assertEqual(by_branch["077"]["modelo"], "35 HP")
        self.assertEqual(by_branch["195"]["numero_serie"], "10231124")
        self.assertEqual(by_branch["241"]["estado_equipo"], "Compartido con Sucursal 239")
        self.assertTrue(all(item.get("notificacion_sucursal_pendiente") for item in items))

    def test_novedades_validas_persisten_foto_auditoria_y_campos_por_tipo(self):
        equipo = self._equipo()
        tecman.save_grupos_electrogenos({"grupos_electrogenos": [equipo]})
        self._sucursal_session(equipo["sucursal"])
        for tipo in tecman.GENERADOR_NOVEDAD_TIPOS:
            response = self._post_novedad(equipo, tipo)
            self.assertEqual(response.status_code, 302)

        saved = tecman.load_grupos_electrogenos()["grupos_electrogenos"][0]
        self.assertEqual([n["tipo"] for n in saved["novedades"]], list(tecman.GENERADOR_NOVEDAD_TIPOS))
        self.assertTrue(all(n["actor"] == equipo["sucursal"] for n in saved["novedades"]))
        self.assertTrue(all(n["cargado_at"] and n["fecha_evento"] for n in saved["novedades"]))
        self.assertEqual(saved["novedades"][0]["proveedor_tecnico"], "")
        self.assertEqual(saved["novedades"][1]["proveedor_tecnico"], "Técnico Test")
        self.assertEqual(saved["novedades"][2]["proveedor_tecnico"], "")
        self.assertEqual(sum(h["accion"] == "novedad_registrada" for h in saved["historial"]), 3)
        for novedad in saved["novedades"]:
            photo = tecman.GRUPOS_ELECTROGENOS_UPLOADS_DIR / novedad["foto"]["archivo"]
            self.assertTrue(photo.is_file())
            self.assertEqual(novedad["foto"]["mime"], "image/png")
            self.assertEqual(novedad["foto"]["sha256"], tecman.hashlib.sha256(self.PNG_1X1).hexdigest())
        notifications = tecman.load_notif_admin()["notificaciones"]
        self.assertEqual(len(notifications), 1)
        self.assertIn(equipo["sucursal"], notifications[0]["titulo"])
        self.assertIn(equipo["id"], notifications[0]["detalle"])
        self.client.get("/suc/grupos-electrogenos")
        self.assertEqual(len(tecman.load_notif_admin()["notificaciones"]), 1)

    def test_novedades_invalidas_no_mutan_ni_dejan_archivos(self):
        equipo = self._equipo()
        tecman.save_grupos_electrogenos({"grupos_electrogenos": [equipo]})
        self._sucursal_session(equipo["sucursal"])
        valid_datetime = (tecman.datetime.datetime.now() - tecman.datetime.timedelta(minutes=1)).replace(second=0, microsecond=0).isoformat(timespec="minutes")
        tomorrow = (tecman.datetime.datetime.now() + tecman.datetime.timedelta(days=1)).replace(second=0, microsecond=0).isoformat(timespec="minutes")
        cases = [
            ({"tipo": "encendido_prueba", "fecha_evento": valid_datetime}, 400),
            ({"tipo": "encendido_prueba", "fecha_evento": "22/09/2026", "foto": (io.BytesIO(self.PNG_1X1), "a.png")}, 400),
            ({"tipo": "encendido_prueba", "fecha_evento": tomorrow, "foto": (io.BytesIO(self.PNG_1X1), "a.png")}, 400),
            ({"tipo": "problema_falla", "fecha_evento": valid_datetime, "observacion": "", "foto": (io.BytesIO(self.PNG_1X1), "a.png")}, 400),
            ({"tipo": "encendido_prueba", "fecha_evento": valid_datetime, "foto": (io.BytesIO(self.PNG_1X1), "a.gif")}, 400),
            ({"tipo": "encendido_prueba", "fecha_evento": valid_datetime, "foto": (io.BytesIO(b"no-es-imagen"), "a.png")}, 400),
            ({"tipo": "encendido_prueba", "fecha_evento": valid_datetime, "foto": (io.BytesIO(self.PNG_1X1), "a.jpg")}, 400),
            ({"_csrf_token": "malo", "tipo": "encendido_prueba", "fecha_evento": valid_datetime, "foto": (io.BytesIO(self.PNG_1X1), "a.png")}, 400),
        ]
        for payload, status in cases:
            with self.subTest(payload={k: v for k, v in payload.items() if k != "foto"}):
                payload.setdefault("_csrf_token", "csrf-test")
                response = self.client.post(
                    f"/suc/grupos-electrogenos/{equipo['id']}/novedades",
                    data=payload,
                    content_type="multipart/form-data",
                )
                self.assertEqual(response.status_code, status)
                response.close()
                current = tecman.load_grupos_electrogenos()["grupos_electrogenos"][0]
                self.assertNotIn("novedades", current)
                self.assertEqual(list(tecman.GRUPOS_ELECTROGENOS_UPLOADS_DIR.iterdir()), [])
        oversized = FileStorage(
            stream=io.BytesIO(b"\x89PNG\r\n\x1a\n" + b"x" * tecman.GENERADOR_FOTO_MAX_BYTES),
            filename="grande.png",
        )
        with self.assertRaisesRegex(ValueError, "supera el límite"):
            tecman._generador_preparar_foto(oversized)
        oversized.close()
        self.assertEqual(tecman.load_notif_admin()["notificaciones"], [])

    def test_foto_protegida_por_referencia_equipo_y_alcance(self):
        one = self._equipo(tecman.SUCURSALES[0], "Visible")
        two = self._equipo(tecman.SUCURSALES[1], "Oculto")
        tecman.save_grupos_electrogenos({"grupos_electrogenos": [one, two]})
        self._sucursal_session(one["sucursal"])
        self.assertEqual(self._post_novedad(one, "encendido_prueba").status_code, 302)
        saved = tecman.load_grupos_electrogenos()["grupos_electrogenos"]
        one_saved = next(x for x in saved if x["id"] == one["id"])
        novedad = one_saved["novedades"][0]
        good = f"/grupos-electrogenos/{one['id']}/novedades/{novedad['id']}/foto"
        wrong_equipment = f"/grupos-electrogenos/{two['id']}/novedades/{novedad['id']}/foto"
        response = self.client.get(good)
        self.assertEqual(response.status_code, 200)
        response.close()
        direct_static = f"/static/uploads/grupos_electrogenos/{novedad['foto']['archivo']}"
        self.assertEqual(self.client.get(direct_static).status_code, 404)
        self.assertEqual(self.client.get(wrong_equipment).status_code, 404)
        self.assertEqual(self.client.get(f"/grupos-electrogenos/{one['id']}/novedades/no-referenciada/foto").status_code, 404)
        self._sucursal_session(two["sucursal"])
        self.assertEqual(self.client.get(good).status_code, 404)
        self._admin_session()
        response = self.client.get(good)
        self.assertEqual(response.status_code, 200)
        response.close()
        with self.client.session_transaction() as sess:
            sess.clear()
        self.assertEqual(self.client.get(good).status_code, 403)

    def test_novedad_cross_equipo_sesion_invalida_y_error_guardado_limpian(self):
        one = self._equipo(tecman.SUCURSALES[0])
        two = self._equipo(tecman.SUCURSALES[1])
        tecman.save_grupos_electrogenos({"grupos_electrogenos": [one, two]})
        self._sucursal_session(one["sucursal"])
        self.assertEqual(self._post_novedad(two, "encendido_prueba").status_code, 404)
        with mock.patch.object(tecman, "_session_auth_is_valid", return_value=False):
            self.assertEqual(self._post_novedad(one, "encendido_prueba").status_code, 403)
        self._sucursal_session(one["sucursal"])
        with mock.patch.object(tecman, "save_grupos_electrogenos", side_effect=OSError("fallo simulado")):
            with self.assertRaises(OSError):
                self._post_novedad(one, "encendido_prueba")
        self.assertEqual(list(tecman.GRUPOS_ELECTROGENOS_UPLOADS_DIR.iterdir()), [])
        saved = tecman.load_grupos_electrogenos()["grupos_electrogenos"]
        self.assertTrue(all(not x.get("novedades") for x in saved))

    def test_admin_filtra_y_renderiza_novedades_sin_edicion(self):
        equipo = self._equipo()
        otro = self._equipo(marca="Sin novedad")
        tecman.save_grupos_electrogenos({"grupos_electrogenos": [equipo, otro]})
        self._sucursal_session(equipo["sucursal"])
        self._post_novedad(equipo, "mantenimiento_proveedor")
        self._admin_session()
        response = self.client.get("/admin/grupos-electrogenos?tipo_novedad=mantenimiento_proveedor")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Equipo Test", response.data)
        self.assertNotIn(b"Sin novedad", response.data)
        self.assertIn("Mantenimiento del proveedor".encode(), response.data)
        self.assertNotIn(b"Borrar novedad", response.data)

    def test_compatibilidad_payload_db_y_json(self):
        equipo = self._equipo()
        equipo["novedades"] = [{"id": "nov-1", "tipo": "encendido_prueba", "foto": {"archivo": "x.png"}}]
        tecman.save_grupos_electrogenos({"grupos_electrogenos": [equipo]})
        self.assertEqual(tecman.load_grupos_electrogenos()["grupos_electrogenos"][0]["novedades"][0]["id"], "nov-1")
        from models import GrupoElectrogenoDB
        row = GrupoElectrogenoDB.from_dict(equipo)
        self.assertEqual(row.to_dict()["novedades"][0]["foto"]["archivo"], "x.png")


if __name__ == "__main__":
    unittest.main()
